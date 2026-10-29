"""#85c: the drop-downs on the Excel export (src/clark/excel.py), read back with openpyxl. Fictional data.

Run: python3 dev/run_tests.py test_excel_lists
Every fixed-value column has a list rule with exactly clark's values; the project's own lists (workstreams, milestones, use
cases) are on a hidden "Lists" sheet that the import ignores; Owner stays free text. The colour rules are in test_excel_format.
"""
import re
import shutil
import tempfile
import unittest
from pathlib import Path

import _paths  # noqa: F401  (puts src/ and vendor/ on the path)
from openpyxl import load_workbook

from clark import excel
from clark import excel_import as xi
from test_excel import Demo, ON
from test_excel_format import action

YES_NO = ["Yes", "No"]
PRIORITY = ["Urgent", "High", "Medium", "Low"]
RATING = ["High", "Medium", "Low"]
WS, MS, UC = "Workstreams", "Milestones", "Use cases"      # the Lists columns, by header

# sheet -> {column label: (values, or the Lists column it is read from; allow_blank)}. Everything not here has no list.
EXPECTED = {
    "Milestones": {"Project": (["ABC"], False), "Workstream": (WS, True), "Status": (["Planned", "Done", "Cancelled"], False),
                   "State": (["On track", "At risk", "Slipped", "Overdue", "Forecast", "Achieved"], True),
                   "State set by you": (YES_NO, False), "Movement": (["hold", "rebaseline", "recover", "escalate"], True),
                   "Client-visible": (YES_NO, False)},
    "Actions": {"Project": (["ABC"], False), "Workstream": (WS, True), "Priority": (PRIORITY, False),
                "Status": (["Open", "Delayed", "Closed", "Cancelled"], False), "Client-visible": (YES_NO, False),
                "Milestone": (MS, True), "Use case": (UC, True)},
    "Risks and issues": {"Project": (["ABC"], False), "Type": (["Risk", "Assumption", "Issue", "Dependency"], False),
             "Workstream": (WS, True), "Impact": (RATING, True), "Likelihood": (RATING, True),
             "Status": (["Open", "Closed", "Cancelled"], False), "Client-visible": (YES_NO, False),
             "Milestone": (MS, True), "Use case": (UC, True)},
    "Decisions": {"Project": (["ABC"], False), "Status": (["Pending", "Made", "Withdrawn"], False),
                  "Client-visible": (YES_NO, False), "Use case": (UC, True)},
    "Use cases": {"Project": (["ABC"], False), "Status": (["On track", "Flag", "High risk"], False),
                  "Client-visible": (YES_NO, False)},
}
# Free text on purpose: people (until the people table, #115), names, notes and dates. Priority on RAID is worked out (locked).
FREE = {"Owner", "Made by", "Description", "Decision", "Milestone", "Due", "Labels"}


def lists_of(wb):
    """{Lists header: [values]} read from the hidden sheet."""
    ws = wb["Lists"]
    out = {}
    for col in ws.iter_cols(min_row=1):
        out[col[0].value] = [c.value for c in col[1:] if c.value is not None]
    return out


def rules_by_label(ws):
    """{header label: DataValidation} for the list rules; each rule is on one column from row 2 to the last spare row."""
    head = {c.column_letter: c.value for c in ws[1]}
    out = {}
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list":
            continue
        (rng,) = str(dv.sqref).split()
        letter = re.match(r"([A-Z]+)2:\1\d+$", rng).group(1)
        out[head[letter]] = dv
    return out


class Lists(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d = Demo()
        cls.d.export()
        cls.wb = load_workbook(cls.d.path)
        cls.lists = lists_of(cls.wb)

    @classmethod
    def tearDownClass(cls):
        cls.d.close()

    def check_sheet(self, name):
        ws = self.wb[name]
        got = rules_by_label(ws)
        self.assertEqual(set(got), set(EXPECTED[name]), name)
        for label, (want, blank) in EXPECTED[name].items():
            dv = got[label]
            with self.subTest(sheet=name, column=label):
                if isinstance(want, list):
                    self.assertEqual(dv.formula1, '"' + ",".join(want) + '"')          # exactly clark's values
                else:
                    m = re.match(r"Lists!\$([A-Z])\$2:\$\1\$(\d+)$", dv.formula1)
                    self.assertIsNotNone(m, dv.formula1)
                    col = self.wb["Lists"][m.group(1) + "1"].value
                    self.assertEqual(col, want)                                        # the right project list
                    self.assertEqual(int(m.group(2)), len(self.lists[want]) + 1)       # and the whole of it
                self.assertEqual(bool(dv.allow_blank), blank)
                self.assertEqual(dv.errorStyle, "stop")                                # anything else is refused
                self.assertTrue(dv.showErrorMessage)
                self.assertEqual(dv.error, "Pick one of the values in the list.")
                self.assertEqual(dv.errorTitle, "Not on the list")
                self.assertFalse(dv.errorTitle.lower().startswith("clark"))

    def test_milestones(self):
        self.check_sheet("Milestones")

    def test_actions(self):
        self.check_sheet("Actions")

    def test_risks_and_issues(self):
        self.check_sheet("Risks and issues")

    def test_decisions(self):
        self.check_sheet("Decisions")

    def test_use_cases(self):
        self.check_sheet("Use cases")

    def test_the_list_values_are_the_stores_own_check_values(self):
        """The inline values are those the database accepts: a value outside them is refused there too."""
        schema = (Path(__file__).resolve().parent.parent / "src" / "clark" / "schema.sql").read_text(encoding="utf-8")
        checks = {("actions", "status"), ("raid", "status"), ("milestones", "status"), ("decisions", "status"),
                  ("use_cases", "status"), ("actions", "priority"), ("raid", "type"), ("raid", "impact"),
                  ("raid", "likelihood"), ("raid", "priority"), ("milestones", "state"), ("milestones", "movement")}
        for table, field in sorted(checks):
            table_sql = schema[schema.index(f"CREATE TABLE IF NOT EXISTS {table} ("):]
            table_sql = table_sql[:table_sql.index(");")]
            m = re.search(rf"\b{field}\s+TEXT[^\n]*?IN \(([^)]*)\)", table_sql)
            self.assertIsNotNone(m, (table, field))
            in_schema = set(re.findall(r"'([^']*)'", m.group(1)))
            with self.subTest(table=table, field=field):
                self.assertEqual(set(excel.choices(table, field)), in_schema)

    def test_the_projects_own_lists_are_on_the_lists_sheet(self):
        self.assertEqual(self.lists[WS], ["Build"])
        self.assertEqual(self.lists[MS], sorted(r["id"] for r in self.d.store.read("milestones", "ABC")))
        self.assertEqual(self.lists[UC], sorted(r["id"] for r in self.d.store.read("use_cases", "ABC")))

    def test_the_lists_sheet_is_hidden_and_last(self):
        self.assertEqual(self.wb["Lists"].sheet_state, "hidden")
        self.assertEqual(self.wb.sheetnames[-1], "Lists")
        self.assertEqual(self.wb.active.title, "Cover")             # the workbook opens on a visible sheet
        for name in self.wb.sheetnames[:-1]:
            self.assertEqual(self.wb[name].sheet_state, "visible", name)

    def test_owner_is_free_text_with_a_note(self):
        for name, label in (("Actions", "Owner"), ("Risks and issues", "Owner"), ("Milestones", "Owner"), ("Decisions", "Made by")):
            ws = self.wb[name]
            self.assertNotIn(label, rules_by_label(ws), name)
            letter = [c.column_letter for c in ws[1] if c.value == label][0]
            notes = [dv for dv in ws.data_validations.dataValidation
                     if dv.type is None and str(dv.sqref).startswith(letter + "2:")]
            self.assertEqual(len(notes), 1, name)
            self.assertTrue(notes[0].showInputMessage)
            self.assertEqual(notes[0].prompt, "Type the person's name.")
            self.assertFalse(notes[0].showErrorMessage)                # nothing is refused

    def test_locked_columns_say_what_to_do_next(self):
        for ws in self.wb:
            for dv in ws.data_validations.dataValidation:
                if dv.type == "custom" and dv.formula1 == "FALSE":
                    want = ("Priority follows Impact and Likelihood. Change those instead."
                            if ws.title == "Risks and issues" and str(dv.sqref).startswith(
                                [c.column_letter for c in ws[1] if c.value == "Priority"][0] + "2:")
                            else "This column is filled in by clark. Leave it as it is.")
                    self.assertEqual(dv.error, want, ws.title)
                    self.assertEqual(dv.errorTitle, "Not editable")

    def test_no_other_column_has_a_list(self):
        for name, wanted in EXPECTED.items():
            for label in rules_by_label(self.wb[name]):
                self.assertNotIn(label, FREE - set(wanted), (name, label))

    def test_texts_a_person_reads_are_plain_words(self):
        for ws in self.wb:
            for dv in ws.data_validations.dataValidation:
                for text in (dv.errorTitle, dv.error, dv.promptTitle, dv.prompt):
                    if text:
                        self.assertFalse(text.lower().startswith("clark"), (ws.title, text))


class ProjectLists(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp(prefix="clark-excel-"))
        self.n = 0

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def build(self, rows):
        self.n += 1
        path = self.dir / f"{self.n}.xlsx"
        excel.export(path, "DEMO", rows, self.dir / f"s{self.n}.json", on=ON)
        return path

    def test_a_long_list_works_because_it_is_on_the_sheet_not_in_the_rule(self):
        """An inline list is limited to 255 characters; 60 long workstream names are not."""
        names = [f"Workstream number {i:02d} with a long name" for i in range(60)]
        wb = load_workbook(self.build({"workstreams": [{"name": n} for n in names], "actions": [action(1)]}))
        self.assertEqual(lists_of(wb)["Workstreams"], sorted(names))
        self.assertLess(len(rules_by_label(wb["Actions"])["Workstream"].formula1), 40)

    def test_workstreams_come_from_the_projects_table_and_skip_archived(self):
        rows = {"workstreams": [{"name": "Build"}, {"name": "Run"}, {"name": "Old", "archived": 1}],
                "actions": [action(1, workstream="Elsewhere")]}
        self.assertEqual(lists_of(load_workbook(self.build(rows)))["Workstreams"], ["Build", "Run"])

    def test_without_a_workstream_table_the_ones_in_use_are_listed(self):
        rows = {"actions": [action(1, workstream="Run"), action(2, workstream="Build"), action(3)]}
        self.assertEqual(lists_of(load_workbook(self.build(rows)))["Workstreams"], ["Build", "Run"])

    def test_an_empty_list_has_no_rule_and_no_sheet(self):
        """A list with nothing in it would refuse every value, so there is none (and no Lists sheet if nothing is listed)."""
        wb = load_workbook(self.build({"actions": [action(1)]}))
        self.assertNotIn("Lists", wb.sheetnames)
        labels = set(rules_by_label(wb["Actions"]))
        self.assertTrue({"Workstream", "Milestone", "Use case"}.isdisjoint(labels))
        self.assertIn("Priority", labels)

    def test_a_name_that_looks_like_a_formula_is_text_in_the_list_and_in_the_cell(self):
        rows = {"workstreams": [{"name": "=Build"}], "actions": [action(1, workstream="=Build")]}
        wb = load_workbook(self.build(rows))
        cell = wb["Actions"]["D2"].value
        self.assertEqual(lists_of(wb)["Workstreams"], [cell])          # the choice and the cell agree
        self.assertEqual(wb["Lists"]["A2"].data_type, "s")
        self.assertEqual(excel.read_workbook(self.dir / "1.xlsx")["actions"][0]["workstream"], "=Build")

    def test_the_same_data_gives_the_same_bytes(self):
        rows = {"workstreams": [{"name": "Build"}], "actions": [action(1, priority="High", due_date="2030-01-01")]}
        self.assertEqual(self.build(rows).read_bytes(), self.build(rows).read_bytes())


class ImportIgnoresLists(unittest.TestCase):
    def setUp(self):
        self.d = Demo()
        self.d.export()

    def tearDown(self):
        self.d.close()

    def test_the_import_does_not_see_the_lists_sheet(self):
        self.assertNotIn("lists", xi.SHEET_TABLES)
        p = xi.read_workbook(self.d.store, self.d.path, "ABC", state_path=self.d.state)
        self.assertEqual(p.problems, [])
        self.assertEqual([n for n in p.notes if "Lists" in n or "ignored" in n], [])

    def test_the_export_and_the_import_with_no_edits_give_no_changes(self):
        p = xi.read_workbook(self.d.store, self.d.path, "ABC", state_path=self.d.state)
        self.assertEqual((p.changes, p.problems), ([], []))
        self.assertGreater(p.unchanged, 0)

    def test_a_change_to_the_lists_sheet_is_not_a_change_to_the_data(self):
        wb = load_workbook(self.d.path)
        wb["Lists"]["A2"].value = "Something else"
        wb.save(self.d.path)
        p = xi.read_workbook(self.d.store, self.d.path, "ABC", state_path=self.d.state)
        self.assertEqual((p.changes, p.problems), ([], []))

    def test_the_reader_of_the_workbook_has_the_five_tables_only(self):
        self.assertEqual(set(excel.read_workbook(self.d.path)), set(excel.TABLES))


if __name__ == "__main__":
    unittest.main()
