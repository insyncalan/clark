"""#38: the round trip. A workbook written by the export (#37, src/clark/excel.py) is edited with openpyxl, as a person would edit
it in Excel, read by the import (src/clark/excel_import.py), applied as one step, and undone; the export's own guard
(edits not yet imported) lets go once the import has run.

Run: python3 dev/run_tests.py test_excel_roundtrip
"""
import json
import shutil
import zipfile
import tempfile
import unittest
from datetime import date
from pathlib import Path

import _paths  # noqa: F401  (puts src/ and vendor/ on the path)
from _demo import USER, DemoStore
from openpyxl import load_workbook

from clark import excel
from clark import excel_import as xi
from clark import page
from clark.store import columns, open_store


_NO_BRAND = tempfile.mkdtemp(prefix="clark-no-brand-")    # a brand logo on this computer never changes a test (#133)


def setUpModule():
    global _OLD_BRANDING
    _OLD_BRANDING, page.BRANDING = page.BRANDING, Path(_NO_BRAND)


def tearDownModule():
    page.BRANDING = _OLD_BRANDING
    shutil.rmtree(_NO_BRAND, ignore_errors=True)


ON = "2030-06-01"
TABLES = ("actions", "raid", "decisions", "milestones", "use_cases", "item_labels")


def edit(path, sheet, item_id, **cells):
    """Set cells of the row with this ID, by the header LABEL the export writes (None clears a cell), as Excel would."""
    wb = load_workbook(str(path))
    ws = wb[sheet]
    head = {c.value: c.column for c in ws[1]}
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, head["ID"]).value == item_id)
    for label, value in cells.items():
        ws.cell(row, head[label]).value = value
    wb.save(str(path))


def add_row(path, sheet, **cells):
    """Append a row under the data, by header LABEL, with no ID (a new item), as a person typing under the table would."""
    wb = load_workbook(str(path))
    ws = wb[sheet]
    row = max(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value is not None) + 1
    head = {c.value: c.column for c in ws[1]}
    for label, value in cells.items():
        ws.cell(row, head[label]).value = value
    wb.save(str(path))


def delete_row(path, sheet, item_id):
    wb = load_workbook(str(path))
    ws = wb[sheet]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == item_id)
    ws.delete_rows(row)
    wb.save(str(path))


class RoundTrip(unittest.TestCase):
    def setUp(self):
        self.demo = DemoStore()
        self.s = self.demo.store
        self.dir = Path(tempfile.mkdtemp(prefix="clark-roundtrip-"))
        self.path = self.dir / "ABC work items.xlsx"
        self.state = self.dir / "state.json"
        with self.s.step("dashboard", USER) as st:
            self.a1 = st.create("actions", {
                "project": "ABC", "description": "Draft the cut-over plan", "workstream": "Build", "owner": "Demo Owner",
                "priority": "High", "due_date": "2030-04-15", "raised": "2030-01-10", "latest_update": "First draft shared",
                "client_message": "The plan is being drafted", "milestone": self.demo.milestone,
                "use_case": self.demo.use_case, "labels": "Sensitive", "source": "demo"})
            self.a2 = st.create("actions", {"project": "ABC", "description": "Book the demo room", "raised": "2026-10-01",
                                            "client_visible": 0, "source": "demo"})
            self.r1 = st.create("raid", {"project": "ABC", "type": "Risk", "description": "Test data arrives late",
                                         "impact": "High", "likelihood": "Medium", "due_date": "2030-05-01",
                                         "response_action": "Ask for a sample early", "raised": "2030-01-12", "source": "demo"})
            self.d1 = st.create("decisions", {"project": "ABC", "decision": "Choose the second team", "status": "Pending",
                                              "needed_by": "2030-07-01", "source": "demo"})
            self.m2 = st.create("milestones", {"project": "ABC", "name": "Slipped milestone", "baseline_date": "2030-08-01",
                                               "forecast_date": "2030-08-20", "percent_complete": 40, "owner": "Demo Owner",
                                               "internal_note": "Keep an eye", "source": "demo"})
        # the demo milestone is rebaselined, so its Original column differs from its Baseline (and a decision is recorded)
        self.demo.update("milestones", self.demo.milestone, {"baseline_date": "2030-04-01"},
                         rebaseline={"agreed_with": "Demo Sponsor", "reason": "Scope grew", "source": "Steering, 01-02-2030"})
        excel.export_project(self.s, "ABC", self.path, self.state, on=ON)

    def tearDown(self):
        self.demo.close()
        shutil.rmtree(self.dir, ignore_errors=True)

    def read(self):
        return xi.read_workbook(self.s, self.path, "ABC", state_path=self.state)

    def steps(self):
        return self.s.con.execute("SELECT count(*) FROM applies").fetchone()[0]

    def snapshot(self):
        """Every table the import can touch, without updated_at (an Undo restores values, not the time it was undone)."""
        return {t: sorted(({k: v for k, v in r.items() if k != "updated_at"} for r in self.s.read(t)),
                          key=lambda r: json.dumps(r, sort_keys=True, default=str)) for t in TABLES}

    def export_again(self):
        return excel.export_project(self.s, "ABC", self.path, self.state, on=ON)

    def edited(self):
        """The usual few edits: an action's owner and due date, a decision's status, a cleared nullable cell."""
        edit(self.path, "Actions", self.a1, Owner="Sam Example", Due=date(2030, 5, 20), **{"Latest update": None})
        edit(self.path, "Decisions", self.d1, Status="Withdrawn")


class FormulaGuard(unittest.TestCase):
    """Text that could be a formula is written as a text cell, as typed; the import must give it back unchanged."""

    TEXTS = ("=1+1", "+1", "-cmd", "@SUM(A1)", "'quoted", "'=1", "a=b", "-5", " =1", "x ", "  padded  ", " 'q")

    def test_text_that_looks_like_a_formula_round_trips_without_a_change(self):
        demo = DemoStore()
        tmp = Path(tempfile.mkdtemp(prefix="clark-formula-"))
        self.addCleanup(shutil.rmtree, tmp, True)
        self.addCleanup(demo.close)
        with demo.store.step("dashboard", USER) as st:
            ids = [st.create("actions", {"project": "ABC", "description": text, "raised": "2026-10-01", "source": "demo"})
                   for text in self.TEXTS]
        path, state = tmp / "ABC work items.xlsx", tmp / "state.json"
        excel.export_project(demo.store, "ABC", path, state, on=ON)
        plan = xi.read_workbook(demo.store, path, "ABC")
        self.assertEqual([c for c in plan.changes if c.kind != "unchanged"], [], "no false change from the apostrophes")
        edit(path, "Actions", ids[0], Owner="Sam Example")                # an ordinary edit still imports, text intact
        plan = xi.read_workbook(demo.store, path, "ABC")
        plan.accept_all()
        xi.apply_import(demo.store, plan, USER, state_path=state)
        got = {r["id"]: r["description"] for r in demo.store.read("actions", "ABC") if r["id"] in ids}
        self.assertEqual([got[i] for i in ids], list(self.TEXTS))


class Reading(RoundTrip):
    def test_every_exported_label_reads_back_as_a_column_a_derived_field_or_labels(self):
        for table, labels in excel.COLUMN_OF.items():
            for label, f in labels.items():
                self.assertTrue(f in columns(self.s.con, table) or f in excel.DERIVED or f == xi.LABELS, (table, label, f))
        self.assertEqual({t: excel.SHEET_OF[t] for t in excel.TABLES},
                         {"milestones": "Milestones", "actions": "Actions", "raid": "Risks and issues", "decisions": "Decisions",
                          "use_cases": "Use cases"})

    def test_an_unchanged_export_imports_as_no_changes_and_nothing_is_ignored(self):
        p = self.read()
        self.assertEqual((p.changes, p.problems), ([], []))
        self.assertEqual([n for n in p.notes if "ignored" in n], [])
        self.assertEqual(p.unchanged, 8)      # 2 actions, 1 risk, 2 decisions (one from the rebaseline), 2 milestones, 1 use case

    def test_the_two_worked_out_columns_are_never_read_back(self):
        edit(self.path, "Milestones", self.m2, Original=date(2031, 1, 1), **{"Shown state": "Achieved"})
        p = self.read()
        self.assertEqual((p.changes, p.problems), ([], []))

    def test_labels_flags_and_dates_typed_as_text_are_read(self):
        edit(self.path, "Actions", self.a1, Due="20-05-2030", **{"Client-visible": "No", "Labels": None})
        (c,) = self.read().changes
        self.assertEqual(c.after, {"due_date": "2030-05-20", "client_visible": 0, "labels": None})
        edit(self.path, "Actions", self.a1, Labels="sensitive")          # the same label in another case: no change on its own
        (c,) = self.read().changes
        self.assertNotIn("labels", c.after)

    def test_a_milestone_agreed_date_edit_in_the_workbook_is_ignored_with_a_note(self):
        edit(self.path, "Milestones", self.m2, Agreed=date(2030, 9, 1), Owner="Sam Example")
        p = self.read()
        (c,) = p.changes
        self.assertEqual(c.after, {"owner": "Sam Example"})
        self.assertTrue(any("rebaseline" in n for n in p.notes))

    def test_database_column_names_still_work_as_headers(self):
        wb = load_workbook(str(self.path))
        ws = wb["Actions"]
        for c in ws[1]:
            c.value = {"ID": "id", "Owner": "owner"}.get(c.value, c.value)
        wb.save(str(self.path))
        edit_wb = load_workbook(str(self.path))
        edit_wb["Actions"].cell(2, 5).value = "Sam Example"
        edit_wb.save(str(self.path))
        (c,) = self.read().changes
        self.assertEqual((c.id, c.after), (self.a1, {"owner": "Sam Example"}))


class Importing(RoundTrip):
    def test_edit_import_and_undo_restores_the_store_in_one_step(self):
        before, n = self.snapshot(), self.steps()
        self.edited()
        p = self.read()
        self.assertEqual(p.problems, [])
        self.assertEqual({(c.id, tuple(c.after)) for c in p.changes},
                         {(self.a1, ("owner", "due_date", "latest_update")), (self.d1, ("status",))})
        p.accept_all()
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertTrue(res.applied, res.message)
        self.assertEqual(res.counts, {"new": 0, "changed": 2, "removed": 0})
        a = self.s.read("actions", ref=self.a1)
        self.assertEqual((a["owner"], a["due_date"], a["latest_update"]), ("Sam Example", "2030-05-20", None))
        self.assertEqual(self.s.read("decisions", ref=self.d1)["status"], "Withdrawn")
        self.assertEqual(self.steps(), n + 1)
        self.assertEqual(res.step_id, self.demo.last_step())
        out = self.s.undo(res.step_id, USER)
        self.assertEqual(out["skipped"], [])
        self.assertEqual(self.snapshot(), before)

    def test_an_edited_dropdown_value_imports(self):
        """A priority and a status changed to another listed value (as picking from the drop-down does) are staged, and only those."""
        edit(self.path, "Actions", self.a1, Priority="Low")
        edit(self.path, "Decisions", self.d1, Status="Withdrawn")
        p = self.read()
        self.assertEqual(p.problems, [])
        self.assertEqual({(c.id, tuple(c.after)) for c in p.changes},
                         {(self.a1, ("priority",)), (self.d1, ("status",))})

    def test_a_labels_edit_is_applied_and_undone_too(self):
        before = self.snapshot()
        edit(self.path, "Actions", self.a1, Labels=None)
        p = self.read()
        p.accept_all()
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertTrue(res.applied, res.message)
        self.assertEqual(self.s.read("item_labels"), [])
        self.s.undo(res.step_id, USER)
        self.assertEqual(self.snapshot(), before)

    def test_a_row_deleted_in_the_sheet_is_staged_as_removed_and_applied_as_archived(self):
        delete_row(self.path, "Actions", self.a2)
        p = self.read()
        (c,) = p.changes
        self.assertEqual((c.kind, c.id), ("removed", self.a2))
        p.accept_all()
        self.assertEqual(p.accepted(), [])                                               # never accepted in bulk
        p.accept(c.n)
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertTrue(res.applied, res.message)
        row = self.s.read("actions", ref=self.a2)
        self.assertEqual((row["status"], row["archived"]), ("Open", 1))
        self.s.undo(res.step_id, USER)
        self.assertEqual(self.s.read("actions", ref=self.a2)["status"], "Open")

    def test_a_new_row_typed_under_the_data_is_added(self):
        wb = load_workbook(str(self.path))
        ws = wb["Actions"]
        head = {c.value: c.column for c in ws[1]}
        row = ws.max_row + 1
        ws.cell(row, head["Description"]).value = "Order the laptops"
        ws.cell(row, head["Owner"]).value = "Sam Example"
        wb.save(str(self.path))
        p = self.read()
        (c,) = p.changes
        self.assertEqual((c.kind, c.table), ("new", "actions"))
        p.accept_all()
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertTrue(res.applied, res.message)
        (new,) = [a for a in self.s.read("actions", project="ABC") if a["description"] == "Order the laptops"]
        self.assertEqual((new["owner"], new["id"] not in (self.a1, self.a2)), ("Sam Example", True))

    def test_a_bad_value_is_reported_with_the_id_and_field_and_nothing_is_applied(self):
        edit(self.path, "Actions", self.a1, Owner="Sam Example")                       # fine
        edit(self.path, "Actions", self.a2, Status="Finished")                         # not a status
        edit(self.path, "Decisions", self.d1, Status="Maybe")                          # nor this
        p = self.read()
        self.assertEqual({c.id for c in p.changes if c.problem}, {self.a2, self.d1})
        p.accept(*[c.n for c in p.changes])                                            # accepted, bad rows too
        before, n = self.snapshot(), self.steps()
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertFalse(res.applied)
        self.assertEqual({(x.id, x.field) for x in res.problems}, {(self.a2, "status"), (self.d1, "status")})
        self.assertIn(self.a2, res.message)
        self.assertEqual((self.snapshot(), self.steps()), (before, n))
        self.assertFalse(res.marked)

    def test_a_date_that_is_not_a_date_is_reported_with_the_id_and_field(self):
        edit(self.path, "Actions", self.a1, Due="next week")
        p = self.read()
        self.assertEqual(p.changes, [])
        self.assertEqual([(x.id, x.field) for x in p.problems], [(self.a1, "due_date")])


class Guard(RoundTrip):
    def test_exporting_before_importing_is_refused_and_after_importing_is_allowed(self):
        self.edited()
        held = excel.file_hash(self.path)
        with self.assertRaises(excel.Refused) as cm:
            self.export_again()
        self.assertIn("Import it first", str(cm.exception))
        self.assertEqual(excel.file_hash(self.path), held)                              # the edited workbook is untouched
        p = self.read()
        p.accept_all()
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertTrue(res.applied and res.marked, res.message)
        self.export_again()                                                              # now allowed
        self.assertNotEqual(excel.file_hash(self.path), held)
        self.assertEqual(self.read().changes, [])                                        # and the new workbook matches the store
        self.assertEqual(excel.load_state(self.state, self.path)["imported_hash"], None)  # a fresh export starts a fresh record

    def test_the_default_state_file_beside_the_store_is_used_when_none_is_given(self):
        path = excel.export_project(self.s, "ABC", on=ON)                                # default path and default state file
        edit(path, "Actions", self.a1, Owner="Sam Example")
        with self.assertRaises(excel.Refused):
            excel.export_project(self.s, "ABC", on=ON)
        p = xi.read_workbook(self.s, path, "ABC")
        p.accept_all()
        res = xi.apply_import(self.s, p, USER)
        self.assertTrue(res.marked, res.message)
        self.assertTrue((Path(self.s.path).parent / excel.STATE_FILE).exists())
        excel.export_project(self.s, "ABC", on=ON)

    def test_nothing_is_marked_when_the_import_is_rejected_cancelled_or_rolled_back(self):
        self.edited()
        # rejected: nothing accepted
        p = self.read()
        p.reject_all()
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertFalse(res.applied or res.marked)
        # rolled back: a bad accepted row
        edit(self.path, "Actions", self.a2, Status="Finished")
        p = self.read()
        p.accept(*[c.n for c in p.changes])
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertFalse(res.applied or res.marked)
        self.assertIsNone(excel.load_state(self.state, self.path)["imported_hash"])
        with self.assertRaises(excel.Refused):
            self.export_again()

    def test_a_workbook_saved_again_after_it_was_read_is_not_marked_as_imported(self):
        self.edited()
        p = self.read()
        p.accept_all()
        edit(self.path, "Actions", self.a2, Owner="Alex Example")                        # saved after the plan was made
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertTrue(res.applied)
        self.assertFalse(res.marked)
        self.assertIn("saved again", res.message)
        with self.assertRaises(excel.Refused):
            self.export_again()

    def test_a_workbook_with_no_export_record_is_unrecorded_and_is_only_applied_once_confirmed(self):
        self.state.unlink()                                                              # the export's record is gone
        self.edited()
        before = self.snapshot()
        p = self.read()
        self.assertTrue(p.unrecorded and p.changes)
        self.assertTrue(any("no record of exporting" in n for n in p.notes))
        self.assertTrue(any("no longer in the workbook are offered for archiving (nothing is deleted)." in n and "count as removed" not in n for n in p.notes))
        p.accept_all()                                                                   # accepts nothing, silently or not
        self.assertEqual(p.accepted(), [])
        p.accept(p.changes[0].n)                                                         # even by number: apply refuses
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertFalse(res.applied)
        self.assertIn("no record of exporting", res.message)
        for words in (res.message, *[n for n in p.notes if "no record of exporting" in n],
                      *[l.strip() for l in p.report().splitlines() if "no record of exporting" in l]):
            self.assertTrue(words.removeprefix("Note: ").startswith("There is no record of exporting"), words)
        self.assertEqual(self.snapshot(), before)
        p.confirm_unrecorded()                                                           # the person confirms
        p.accept_all()
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertTrue(res.applied and res.marked, res.message)
        self.export_again()
        self.state.unlink()
        edit(self.path, "Actions", self.a2, Owner="Alex Example")
        again = xi.read_workbook(self.s, self.path, "ABC", state_path=self.state, allow_unrecorded=True)   # the parameter does it too
        again.accept_all()
        self.assertTrue(xi.apply_import(self.s, again, USER, state_path=self.state).applied)


class Refusal(RoundTrip):
    def test_a_store_that_refuses_to_open_the_step_gives_a_plain_result_and_the_workbook_is_not_marked(self):
        self.edited()
        p = self.read()
        p.accept_all()
        before, n, held = self.snapshot(), self.steps(), excel.load_state(self.state, self.path)

        class Refusing:                      # a store whose step() is refused on entry (as a failed backup will be, #86)
            path = self.s.path

            def __init__(self, real):
                self.real = real

            def step(self, *a, **k):
                raise xi.Invalid("The backup could not be made, so nothing was changed.")

            def __getattr__(self, name):
                return getattr(self.real, name)

        marked = []
        real_mark = excel.mark_imported
        excel.mark_imported = lambda *a, **k: marked.append(a)
        try:
            res = xi.apply_import(Refusing(self.s), p, USER, state_path=self.state)
        finally:
            excel.mark_imported = real_mark
        self.assertFalse(res.applied or res.marked)
        self.assertIsNone(res.step_id)
        self.assertIn("The backup could not be made", res.message)
        self.assertEqual(marked, [])
        self.assertEqual((self.snapshot(), self.steps()), (before, n))
        self.assertEqual(excel.load_state(self.state, self.path), held)
        with self.assertRaises(excel.Refused):                    # still counts as having unimported edits
            self.export_again()


class Idempotent(RoundTrip):
    """Importing the same file twice must change nothing the second time (review HIGH)."""

    def test_a_second_import_of_the_same_file_changes_and_cancels_nothing(self):
        add_row(self.path, "Actions", Description="Brand new")
        p = self.read()
        p.accept_all()
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertTrue(res.applied and res.marked, res.message)
        self.assertEqual(res.counts, {"new": 1, "changed": 0, "removed": 0})
        new = [a for a in self.s.read("actions", "ABC") if a["description"] == "Brand new"]
        self.assertEqual(len(new), 1)
        before, n = self.snapshot(), self.steps()
        again = self.read()                                                              # the same file, imported again
        self.assertTrue(again.already_imported)
        self.assertEqual(again.changes, [])
        self.assertTrue(any("already imported" in x for x in again.notes))
        again.accept_all()
        res = xi.apply_import(self.s, again, USER, state_path=self.state)
        self.assertFalse(res.applied)
        self.assertIn("already imported", res.message)
        self.assertEqual((self.snapshot(), self.steps()), (before, n))
        self.assertEqual(self.s.read("actions", ref=new[0]["id"])["status"], "Open")     # the new item was not cancelled

    def test_two_plans_of_one_file_apply_only_once(self):
        add_row(self.path, "Actions", Description="Brand new")
        first, second = self.read(), self.read()
        first.accept_all()
        second.accept_all()
        self.assertTrue(xi.apply_import(self.s, first, USER, state_path=self.state).applied)
        n = self.steps()
        res = xi.apply_import(self.s, second, USER, state_path=self.state)
        self.assertFalse(res.applied)
        self.assertEqual(self.steps(), n)

    def test_the_file_saved_again_does_not_turn_the_item_it_made_into_one_removed_from_the_workbook(self):
        add_row(self.path, "Actions", Description="Brand new")
        p = self.read()
        p.accept_all()
        self.assertTrue(xi.apply_import(self.s, p, USER, state_path=self.state).applied)
        edit(self.path, "Actions", self.a2, Owner="Alex Example")                        # saved again: a different file
        again = self.read()
        self.assertFalse(again.already_imported)
        self.assertEqual([c.kind for c in again.changes if c.kind == "removed"], [])     # the new item was not in the export
        self.assertEqual(sorted(c.kind for c in again.changes), ["changed"])             # the new row (no id) is not made again
        self.assertEqual(again.already_created, 1)


class NotTwice(RoundTrip):
    """A no-id row an import already made is not made again when the same content is read from other bytes (review LOW)."""

    def make(self, **cells):
        add_row(self.path, "Actions", **(cells or {"Description": "Brand new"}))
        p = self.read()
        p.accept_all()
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertTrue(res.applied and res.marked, res.message)
        return res

    def resave(self, to=None):
        """The same content saved again by openpyxl: different bytes (a different file hash)."""
        before = self.path.read_bytes()
        out = Path(to) if to else self.path
        wb = load_workbook(str(self.path))
        wb.properties.title = f"Saved again {len(before)}"            # (a save by another program: other bytes, same cells)
        wb.save(str(out))
        wb.close()
        self.assertNotEqual(out.read_bytes(), before)
        return out

    def count(self):
        return len(self.s.read("actions", "ABC"))

    def add_rows(self, *texts, **more):
        """Several new rows (no ID) under the Actions table in one save."""
        wb = load_workbook(str(self.path))
        ws = wb["Actions"]
        head = {c.value: c.column for c in ws[1]}
        row = ws.max_row
        for text in texts:
            row += 1
            ws.cell(row, head["Description"]).value = text
            for label, value in more.items():
                ws.cell(row, head[label]).value = value
        wb.save(str(self.path))

    def test_a_resaved_copy_does_not_make_the_new_row_again_and_cancels_nothing(self):
        res = self.make()
        new_id = res.created[0]["id"]
        self.assertEqual(len(res.created), 1)
        n_rows, n_steps = self.count(), self.steps()
        self.resave()
        again = self.read()
        self.assertFalse(again.already_imported)                    # other bytes: the file is read again...
        self.assertEqual(again.changes, [])                         # ...but nothing is staged
        self.assertEqual(again.already_created, 1)
        self.assertTrue(any(f"Already imported as {new_id}: left as it is." in x for x in again.notes), again.notes)
        self.assertIn("1 already imported", again.report())
        again.accept_all()
        self.assertFalse(xi.apply_import(self.s, again, USER, state_path=self.state).applied)
        self.assertEqual((self.count(), self.steps()), (n_rows, n_steps))
        self.assertEqual(self.s.read("actions", ref=new_id)["status"], "Open")

    def test_a_copy_saved_to_another_path_is_covered_too(self):
        new_id = self.make().created[0]["id"]
        n_rows = self.count()
        copy = self.resave(self.dir / "ABC copy.xlsx")
        plan = xi.read_workbook(self.s, copy, "ABC", state_path=self.state)
        self.assertEqual([c for c in plan.changes if c.kind == "new"], [])
        self.assertTrue(any(f"Already imported as {new_id}" in x for x in plan.notes))
        self.assertEqual(self.count(), n_rows)

    def test_editing_the_new_row_after_importing_it_makes_a_different_new_row(self):
        self.make()
        self.resave()
        wb = load_workbook(str(self.path))
        ws = wb["Actions"]
        head = {c.value: c.column for c in ws[1]}
        row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, head["ID"]).value is None and ws.cell(r, head["Description"]).value)
        ws.cell(row, head["Description"]).value = "Brand new, reworded"
        wb.save(str(self.path))
        p = self.read()
        self.assertEqual([c.kind for c in p.changes], ["new"])
        self.assertEqual(p.already_created, 0)

    def test_two_identical_rows_in_one_import_are_both_made_and_both_recorded(self):
        self.add_rows("Same text", "Same text")
        p = self.read()
        self.assertEqual([c.kind for c in p.changes], ["new", "new"])
        p.accept_all()
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertEqual(res.counts["new"], 2)
        self.assertEqual(len({c["id"] for c in res.created}), 2)
        self.assertEqual(len(excel.load_state(self.state, self.path)["created"]), 2)
        n_rows = self.count()
        self.resave()
        again = self.read()
        self.assertEqual(again.changes, [])                         # each recorded row covers one workbook row
        self.assertEqual(again.already_created, 2)
        add_row(self.path, "Actions", Description="Same text")      # a third identical row typed later is a new row
        third = self.read()
        self.assertEqual([c.kind for c in third.changes], ["new"])
        self.assertEqual(third.already_created, 2)
        self.assertEqual(self.count(), n_rows)

    def test_a_rejected_new_row_is_not_recorded_and_is_staged_again(self):
        self.add_rows("Not wanted", "Wanted")
        p = self.read()
        wanted = next(c for c in p.changes if c.after["description"] == "Wanted")
        p.accept(wanted.n)
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertEqual(len(res.created), 1)
        self.resave()
        again = self.read()
        self.assertEqual([c.after["description"] for c in again.changes], ["Not wanted"])

    def test_nothing_is_recorded_for_a_refused_or_rolled_back_import(self):
        self.add_rows("Fine")
        edit(self.path, "Actions", self.a2, Status="Finished")                  # a bad accepted change: all is rolled back
        p = self.read()
        p.accept(*[c.n for c in p.changes])
        held = excel.load_state(self.state, self.path)
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertFalse(res.applied)
        self.assertEqual(res.created, [])
        self.assertEqual(excel.load_state(self.state, self.path), held)
        self.assertNotIn("created", held)

    def test_nothing_is_recorded_when_the_store_refuses_to_open_the_step(self):
        add_row(self.path, "Actions", Description="Fine")
        p = self.read()
        p.accept_all()

        class Refusing:
            path = self.s.path

            def __init__(self, real):
                self.real = real

            def step(self, *a, **k):
                raise xi.Invalid("The backup could not be made, so nothing was changed.")

            def __getattr__(self, name):
                return getattr(self.real, name)

        held = excel.load_state(self.state, self.path)
        self.assertFalse(xi.apply_import(Refusing(self.s), p, USER, state_path=self.state).applied)
        self.assertEqual(excel.load_state(self.state, self.path), held)

    def test_after_an_undo_the_same_file_makes_the_row_again_and_an_export_starts_a_fresh_record(self):
        res = self.make()
        self.s.undo(res.step_id, USER)
        self.assertEqual(len(excel.load_state(self.state, self.path)["created"]), 1)         # (the record stays; it is not in force)
        self.resave()
        again = self.read()
        self.assertEqual([c.kind for c in again.changes if c.kind == "new"], ["new"])        # the undo was on purpose: made again
        self.assertEqual(again.already_created, 0)
        self.path.unlink()
        self.export_again()                                                                  # an export starts a fresh record
        self.assertNotIn("created", excel.load_state(self.state, self.path))

    def test_a_workbook_saved_again_after_it_was_read_still_records_the_rows_it_made(self):
        add_row(self.path, "Actions", Description="Brand new")
        p = self.read()
        p.accept_all()
        edit(self.path, "Actions", self.a2, Owner="Alex Example")               # saved again after the plan was read
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertTrue(res.applied and not res.marked)
        self.assertEqual(len(excel.load_state(self.state, self.path)["created"]), 1)
        self.assertEqual([c.kind for c in self.read().changes if c.kind == "new"], [])

    def test_the_list_of_fingerprints_is_bounded_and_drops_the_oldest(self):
        rec = {}
        excel.add_created(rec, [{"fp": f"f{i}", "id": f"ABC-A-{i}"} for i in range(excel.MAX_CREATED + 10)])
        self.assertEqual(len(rec["created"]), excel.MAX_CREATED)
        self.assertEqual(rec["created"][0]["fp"], "f10")
        excel.add_created(rec, [{"fp": "last", "id": "x"}])
        self.assertEqual(len(rec["created"]), excel.MAX_CREATED)
        self.assertEqual((rec["created"][0]["fp"], rec["created"][-1]["fp"]), ("f11", "last"))

    def test_the_fingerprint_ignores_the_file_hash_and_column_order(self):
        a = xi.fingerprint("actions", {"description": "x", "owner": "y"})
        self.assertEqual(a, xi.fingerprint("actions", {"owner": "y", "description": "x"}))
        self.assertNotEqual(a, xi.fingerprint("raid", {"description": "x", "owner": "y"}))
        self.assertNotEqual(a, xi.fingerprint("actions", {"description": "x ", "owner": "y"}))

    def test_the_export_stays_byte_deterministic(self):
        first = self.path.read_bytes()
        self.make()
        self.export_again()
        self.export_again()
        again = self.path.read_bytes()
        self.export_again()
        self.assertEqual(again, self.path.read_bytes())
        self.assertNotEqual(first, again)                                        # (the new row is in it now)


class UndoAndRestore(RoundTrip):
    """"""

    def imported(self):
        self.edited()
        p = self.read()
        p.accept_all()
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertTrue(res.applied and res.marked, res.message)
        return res

    def test_import_then_undo_then_the_same_file_can_be_imported_again(self):
        res = self.imported()
        self.assertTrue(self.read().already_imported)                       # in force: refused as already imported
        self.s.undo(res.step_id, USER)
        again = self.read()
        self.assertFalse(again.already_imported)
        self.assertTrue(any("has been undone" in n for n in again.notes), again.notes)
        self.assertTrue(again.changes)
        again.accept(*[c.n for c in again.changes])                         # (the undo made a row newer: a clash, by number)
        res2 = xi.apply_import(self.s, again, USER, state_path=self.state)
        self.assertTrue(res2.applied and res2.marked, res2.message)
        self.assertEqual(self.s.read("actions", ref=self.a1)["owner"], "Sam Example")
        self.assertTrue(self.read().already_imported)                       # and the second import is in force

    def test_import_then_import_again_is_refused_as_already_imported(self):
        self.imported()
        n, before = self.steps(), self.snapshot()
        again = self.read()
        self.assertTrue(again.already_imported)
        again.accept_all()
        res = xi.apply_import(self.s, again, USER, state_path=self.state)
        self.assertFalse(res.applied)
        self.assertIn("already imported", res.message)
        self.assertEqual((self.steps(), self.snapshot()), (n, before))

    def test_a_plan_made_before_an_undo_cannot_hide_behind_the_old_stamp_and_an_undone_undo_puts_it_back_in_force(self):
        res = self.imported()
        undo = self.s.undo(res.step_id, USER)
        self.assertFalse(self.read().already_imported)
        self.s.undo(undo["step_id"], USER)                                  # the undo undone: the import is in force again
        self.assertTrue(self.read().already_imported)

    def restored_store(self):
        """A second store on a copy of the database taken now (before the import)."""
        self.s.con.execute("PRAGMA wal_checkpoint(TRUNCATE)")
        folder = self.dir / "restored"
        folder.mkdir()
        shutil.copy(self.s.path, folder / "clark.db")
        return open_store(folder / "clark.db", onedrive_roots=[])

    def test_a_restored_older_database_can_import_the_file_again(self):
        self.edited()
        old = self.restored_store()
        self.addCleanup(old.close)
        p = self.read()
        p.accept_all()
        self.assertTrue(xi.apply_import(self.s, p, USER, state_path=self.state).applied)
        self.assertTrue(self.read().already_imported)
        with old.step("dashboard", USER) as st:                             # the restored copy goes on, and reuses the step number
            st.update("actions", self.a2, {"owner": "Someone Else"})
        again = xi.read_workbook(old, self.path, "ABC", state_path=self.state)
        self.assertFalse(again.already_imported)
        self.assertTrue(again.changes)
        again.accept(*[c.n for c in again.changes])
        res = xi.apply_import(old, again, USER, state_path=self.state)
        self.assertTrue(res.applied, res.message)
        self.assertEqual(old.read("actions", ref=self.a1)["owner"], "Sam Example")

    def test_a_fingerprint_counts_only_while_its_step_is_in_force(self):
        add_row(self.path, "Actions", Description="Brand new")
        old = self.restored_store()
        self.addCleanup(old.close)
        p = self.read()
        p.accept_all()
        res = xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertTrue(res.applied, res.message)
        edit(self.path, "Actions", self.a2, Owner="Alex Example")            # other bytes
        self.assertEqual(self.read().already_created, 1)                     # in force: not made again
        self.s.undo(res.step_id, USER)
        undone = self.read()
        self.assertEqual(undone.already_created, 0)                          # undone: made again
        self.assertEqual([c.kind for c in undone.changes if c.kind == "new"], ["new"])
        undone.accept(*[c.n for c in undone.changes if c.kind == "new"])
        self.assertTrue(xi.apply_import(self.s, undone, USER, state_path=self.state).applied)
        rows = [a for a in self.s.read("actions", "ABC") if a["description"] == "Brand new"]
        self.assertEqual(len(rows), 1)                                       # (the undo removed the first; this is the second)
        in_old = xi.read_workbook(old, self.path, "ABC", state_path=self.state)       # a restored copy: no such step either
        self.assertEqual(in_old.already_created, 0)


class Wording(RoundTrip):
    def test_a_database_error_is_put_in_plain_words(self):
        self.edited()
        p = self.read()
        p.accept_all()

        class Locked:
            path = self.s.path

            def __init__(self, real):
                self.real = real

            def step(self, *a, **k):
                raise xi.sqlite3.OperationalError("database is locked")

            def __getattr__(self, name):
                return getattr(self.real, name)

        res = xi.apply_import(Locked(self.s), p, USER, state_path=self.state)
        self.assertFalse(res.applied)
        self.assertIn("in use by another program", res.message)
        self.assertNotIn("database is locked", res.message)


class UseCaseDone(RoundTrip):
    """#412: Done is a use case status. Chosen in the workbook it imports as Done; a Done use case exports as Done, in grey, and an
    unchanged workbook imports as no change at all."""

    def test_done_typed_in_the_workbook_imports_as_done_and_undoes(self):
        edit(self.path, "Use cases", self.demo.use_case, Status="Done")
        p = self.read()
        self.assertEqual(p.problems, [])
        p.accept_all()
        self.assertTrue(xi.apply_import(self.s, p, USER, state_path=self.state).applied)
        self.assertEqual(self.s.read("use_cases", ref=self.demo.use_case)["status"], "Done")

    def test_a_done_use_case_exports_as_done_in_grey_and_reads_back_unchanged(self):
        from clark import excel_colours as colours
        with self.s.step("dashboard", USER) as st:
            st.update("use_cases", self.demo.use_case, {"status": "Done"})
        self.export_again()
        ws = load_workbook(str(self.path))["Use cases"]
        head = {c.value: c.column for c in ws[1]}
        row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, head["ID"]).value == self.demo.use_case)
        cell = ws.cell(row, head["Status"])
        self.assertEqual(cell.value, "Done")
        rules = [r for cf in ws.conditional_formatting for r in cf.rules
                 if str(cf.sqref).startswith(cell.column_letter) and r.type == "cellIs" and r.formula == ['"Done"']]
        self.assertEqual(len(rules), 1, "one rule draws Done")
        self.assertEqual((rules[0].dxf.fill.bgColor.rgb[-6:], rules[0].dxf.font.color.rgb[-6:]), colours.hexes("grey"))
        p = self.read()
        self.assertEqual((p.changes, p.problems), ([], []))
        self.assertEqual(xi.read_workbook(self.s, self.path, "ABC", state_path=self.state).problems, [])

    def test_a_status_that_is_not_one_of_the_four_is_still_refused(self):
        edit(self.path, "Use cases", self.demo.use_case, Status="Finished")
        p = self.read()
        self.assertEqual([c.problem for c in p.changes], ["Status must be one of On track, Flag, High risk, Done (got 'Finished')"])
        p.accept_all()
        xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertEqual(self.s.read("use_cases", ref=self.demo.use_case)["status"], "On track", "a change with a problem is not applied")


class RichText(RoundTrip):
    HOSTILE = '<p onerror="steal()">Live</p><script>alert(1)</script><img src=x onerror=alert(2)>'

    def import_standing(self, text):
        edit(self.path, "Use cases", self.demo.use_case, **{"Where it stands": text})
        p = self.read()
        p.accept_all()
        self.assertTrue(xi.apply_import(self.s, p, USER, state_path=self.state).applied)
        return self.s.read("use_cases", ref=self.demo.use_case)["standing"]

    def test_standing_goes_through_the_clean_rich_seam(self):
        seen = []
        real = xi.clean_rich
        xi.clean_rich = lambda field, text: seen.append(field) or "[cleaned] " + text
        try:
            stored = self.import_standing("Going well")
        finally:
            xi.clean_rich = real
        self.assertEqual(seen, ["standing"])
        self.assertEqual(stored, "[cleaned] Going well")
        self.assertEqual(xi.RICH_FIELDS, {"standing"})

    def test_the_seam_is_a_no_op_today_and_hostile_html_is_stored_as_typed(self):
        self.assertEqual(xi.clean_rich("standing", self.HOSTILE), self.HOSTILE)
        self.assertEqual(self.import_standing(self.HOSTILE), xi.clean_rich("standing", self.HOSTILE))

    def test_hostile_html_is_published_without_its_active_part(self):
        try:
            from clark import publish
        except ImportError:
            self.skipTest("clark.publish is not on this branch yet (#39); this test runs once it merges")
        stored = self.import_standing(self.HOSTILE)
        out = publish.clean_html(stored)
        for bad in ("<script", "onerror", "alert("):
            self.assertNotIn(bad, out)
        self.assertIn("Live", out)


class Clash(RoundTrip):
    def test_only_the_cells_the_person_edited_are_staged_and_newer_values_in_clark_are_left_alone(self):
        self.demo.update("actions", self.a1, {"owner": "Kim Example"})                   # changed in clark after the export
        edit(self.path, "Actions", self.a1, Description="Draft the cut-over plan v2")    # the person edits one cell
        p = self.read()
        (c,) = p.changes
        self.assertEqual((c.state, c.clash, c.from_snapshot), ("clash", True, True))
        self.assertEqual(c.after, {"description": "Draft the cut-over plan v2"})
        self.assertEqual(c.kept, {"owner": "Kim Example"})
        text = p.report()
        self.assertIn("You edited Description", text)
        self.assertIn("Newer in clark and left alone: Owner (now Kim Example)", text)
        p.accept(c.n)
        self.assertTrue(xi.apply_import(self.s, p, USER, state_path=self.state).applied)
        row = self.s.read("actions", ref=self.a1)
        self.assertEqual((row["description"], row["owner"]), ("Draft the cut-over plan v2", "Kim Example"))

    def test_a_clash_where_the_person_edits_the_same_cell_stages_it_and_an_untouched_row_is_no_change(self):
        self.demo.update("actions", self.a1, {"owner": "Kim Example"})
        self.demo.update("actions", self.a2, {"owner": "Kim Example"})
        edit(self.path, "Actions", self.a1, Owner="Sam Example")
        p = self.read()
        (c,) = p.changes                                                                  # a2 was only changed in clark: no change
        self.assertEqual((c.after, c.before, c.kept), ({"owner": "Sam Example"}, {"owner": "Kim Example"}, {}))

    def test_without_the_export_snapshot_the_clash_says_unedited_fields_are_also_reverted(self):
        self.demo.update("actions", self.a1, {"owner": "Kim Example"})
        edit(self.path, "Actions", self.a1, Description="Draft the cut-over plan v2")
        state = json.loads(self.state.read_text(encoding="utf-8"))
        for rec in state["workbooks"].values():
            rec["rows"] = None
        self.state.write_text(json.dumps(state), encoding="utf-8")
        (c,) = self.read().changes
        self.assertEqual((c.clash, c.from_snapshot), (True, False))
        self.assertEqual(set(c.after), {"description", "owner"})
        p = self.read()
        self.assertIn("including ones you did not edit", p.report())

    def test_a_row_added_in_clark_after_the_export_is_not_removed_from_the_workbook(self):
        with self.s.step("dashboard", USER) as st:
            late = st.create("actions", {"project": "ABC", "description": "Added later", "raised": "2030-02-01", "source": "demo"})
        p = self.read()
        self.assertEqual(p.changes, [])
        p.accept_all()
        xi.apply_import(self.s, p, USER, state_path=self.state)
        self.assertEqual(self.s.read("actions", ref=late)["status"], "Open")


class TooBig(RoundTrip):
    def bomb(self, name="bomb.xlsx", size=60 * 1024 * 1024, members=1):
        path = self.dir / name
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for i in range(members):
                z.writestr(f"xl/worksheets/sheet{i}.xml", b"\0" * (size // members))
        return path

    def test_a_small_file_that_unpacks_to_a_huge_one_is_refused_in_plain_words_before_it_is_opened(self):
        path = self.bomb()
        self.assertLess(path.stat().st_size, 5 * 1024 * 1024)
        for call in (lambda: xi.read_workbook(self.s, path, "ABC", state_path=self.state), lambda: excel.read_workbook(path)):
            with self.assertRaises((xi.Invalid, excel.Refused)) as cm:
                call()
            self.assertIn("would unpack to", str(cm.exception))
            self.assertIn("bomb.xlsx", str(cm.exception))

    def test_one_member_over_its_limit_or_too_many_members_is_refused(self):
        with self.assertRaises(xi.Invalid):
            xi.read_workbook(self.s, self.bomb("one.xlsx", size=excel.MAX_MEMBER_BYTES + 1), "ABC", state_path=self.state)
        with self.assertRaises(xi.Invalid) as cm:
            xi.read_workbook(self.s, self.bomb("many.xlsx", size=excel.MAX_MEMBERS + 5, members=excel.MAX_MEMBERS + 5), "ABC",
                             state_path=self.state)
        self.assertIn("parts", str(cm.exception))

    def test_a_file_over_the_size_limit_is_refused_and_a_real_export_is_far_below_every_limit(self):
        big = self.dir / "big.xlsx"
        big.write_bytes(b"x" * (excel.MAX_FILE_BYTES + 1))
        with self.assertRaises(xi.Invalid) as cm:
            xi.read_workbook(self.s, big, "ABC", state_path=self.state)
        self.assertIn("more than", str(cm.exception))
        self.assertLess(self.path.stat().st_size, excel.MAX_FILE_BYTES // 100)
        excel.check_workbook_size(self.path)                                             # the real export passes


if __name__ == "__main__":
    unittest.main()
