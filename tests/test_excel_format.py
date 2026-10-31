"""#85a: how the Excel export looks (src/clark/excel.py, src/clark/excel_colours.py), read back with openpyxl. Fictional data.

Run: python3 dev/run_tests.py test_excel_format
#85c: colour is conditional formatting (a rule per word, and a formula rule for overdue dates), so an edited value recolours
itself; no cell holds a formula and the values, drop-downs and the import are not touched by it.
"""
import shutil
import tempfile
import unittest
import zipfile
from datetime import date
from pathlib import Path

import _paths  # noqa: F401  (puts src/ and vendor/ on the path)
from openpyxl import load_workbook

from clark import excel
from clark import excel_colours as colours
from clark import look
from test_excel import Demo, ON

TODAY = date(2030, 6, 1)


def action(n, **kw):
    return {"id": f"DEMO-A-{n:03d}", "project": "DEMO", "description": f"Action {n}", "raised": "2030-01-01",
            "client_message": "", "client_visible": 1, "source": "demo", "updated_at": "2030-01-01 09:30:00.123", **kw}


def raid(n, **kw):
    return {"id": f"DEMO-R-{n:03d}", "project": "DEMO", "type": "Risk", "description": f"Risk {n}", "raised": "2030-01-01",
            "client_message": "", "client_visible": 1, "source": "demo", "updated_at": "2030-01-01 09:30:00.123", **kw}


def decision(n, **kw):
    return {"id": f"DEMO-D-{n:03d}", "project": "DEMO", "decision": f"Decision {n}", "client_message": "",
            "client_visible": 1, "source": "demo", "updated_at": "2030-01-01 09:30:00.123", **kw}


def use_case(n, **kw):
    return {"id": f"DEMO-U-{n:03d}", "project": "DEMO", "title": f"Use case {n}", "ref": "", "headline": "", "standing": "",
            "client_visible": 1, "source": "demo", "updated_at": "2030-01-01 09:30:00.123", **kw}


def rgb(color):
    """A colour's six hex digits, or None when it has none (a default or theme colour)."""
    value = getattr(color, "rgb", None)
    return value[-6:] if isinstance(value, str) else None


def luminance(hexes):
    def part(i):
        v = int(hexes[i:i + 2], 16) / 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    return 0.2126 * part(0) + 0.7152 * part(2) + 0.0722 * part(4)


def contrast(a, b):
    hi, lo = sorted((luminance(a), luminance(b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


def rules_of(ws):
    """[(first column letter, rule)] for every conditional format on a sheet, in priority order."""
    out = []
    for cf in ws.conditional_formatting:
        for r in cf.rules:
            out.append((str(cf.sqref).split(":")[0].rstrip("0123456789"), r))
    return sorted(out, key=lambda x: x[1].priority)


def look_of(ws, col, value):
    """What Excel would draw for `value` in column `col`: (fill hex, text hex) of the rule that matches it, else None. Only the
    equal-to-a-word rules are read; the rules are fixed literals, so this is the same as Excel's own test."""
    for letter, r in rules_of(ws):
        if letter == col and r.type == "cellIs" and r.operator == "equal" and r.formula == ['"' + value + '"']:
            return rgb(r.dxf.fill.bgColor), rgb(r.dxf.font.color)
    return None


class Formatting(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp(prefix="clark-excel-"))
        self.n = 0

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def build(self, rows, on=ON):
        self.n += 1
        path = self.dir / f"{self.n}.xlsx"
        excel.export(path, "DEMO", rows, self.dir / f"state{self.n}.json", on=on)
        return path

    def sheet(self, rows, name):
        return load_workbook(self.build(rows))[name]

    @staticmethod
    def cell(ws, row_id, label):
        head = [c.value for c in ws[1]]
        for r in ws.iter_rows(min_row=2):
            if r[0].value == row_id:
                return r[head.index(label)]
        raise KeyError(row_id)

    def everything(self):
        return {"actions": [action(1, priority="High")], "raid": [raid(1)], "decisions": [decision(1)],
                "use_cases": [use_case(1)], "milestones": []}

    # ── panes, filter, header ────────────────────────────────────────────────

    def test_header_and_the_id_and_title_columns_are_frozen_on_every_sheet(self):
        wb = load_workbook(self.build(self.everything()))
        titles = {"Milestones": "Milestone", "Actions": "Description", "Risks and issues": "Description", "Decisions": "Decision",
                  "Use cases": "Use case"}
        for name, title in titles.items():
            ws = wb[name]
            self.assertEqual(ws.freeze_panes, "C2", name)
            self.assertEqual([ws["A1"].value, ws["B1"].value], ["ID", title], name)

    def test_the_filter_covers_the_header_and_the_data(self):
        rows = {"actions": [action(1), action(2), action(3)], "raid": []}
        wb = load_workbook(self.build(rows))
        self.assertEqual(wb["Actions"].auto_filter.ref, "A1:R4")
        self.assertEqual(wb["Risks and issues"].auto_filter.ref, "A1:U1")
        for name, table, cols, _ in excel.SHEETS:
            self.assertTrue(wb[name].auto_filter.ref.startswith("A1:"), name)

    def test_header_is_bold_white_on_the_looks_brand_navy_and_wraps(self):
        wb = load_workbook(self.build(self.everything()))
        for ws in wb:
            for c in ws[1]:
                self.assertTrue(c.font.b, (ws.title, c.value))
                self.assertEqual(c.font.color.rgb[-6:], look.colour("white"))
                self.assertEqual(c.fill.fgColor.rgb[-6:], look.colour("brand-navy"))
                self.assertTrue(c.alignment.wrap_text)

    def test_the_looks_office_font_on_the_header_and_the_data_cells_of_each_data_sheet(self):
        wb = load_workbook(self.build(self.everything()))
        font = look.office_font()
        for name in ("Milestones", "Actions", "Risks and issues", "Decisions", "Use cases"):
            ws = wb[name]
            self.assertEqual(ws["A1"].font.name, font, (name, "header"))
            self.assertEqual(ws["B2"].font.name, font, (name, "data"))
            self.assertEqual(ws["B2"].font.sz, 11, (name, "data size"))

    # ── widths and wrapping ──────────────────────────────────────────────────

    def test_widths_and_wrap_by_kind(self):
        wb = load_workbook(self.build(self.everything()))
        for name, table, cols, _ in excel.SHEETS:
            ws = wb[name]
            for i, (label, field) in enumerate(cols, start=1):
                letter = ws.cell(row=1, column=i).column_letter
                width = ws.column_dimensions[letter].width
                self.assertGreater(width, 0, (name, label))
                if field == "id":
                    self.assertEqual(width, 10)
                if field in excel.DATES:
                    self.assertEqual(width, 12)
                if field in excel.WIDE:
                    self.assertEqual(width, 48)
        ws = wb["Actions"]
        for label, field in (("Description", True), ("Latest update", True), ("Client wording", True), ("ID", False),
                             ("Owner", False), ("Due", False), ("Priority", False), ("Client-visible", False)):
            c = self.cell(ws, "DEMO-A-001", label)
            self.assertEqual(bool(c.alignment.wrap_text), field, label)
            self.assertEqual(c.alignment.vertical, "top", label)
        self.assertEqual(wb["Milestones"].column_dimensions["B"].width, 40)         # the milestone's name

    def test_row_heights_are_left_automatic(self):
        ws = self.sheet(self.everything(), "Actions")
        self.assertEqual({d.height for d in ws.row_dimensions.values()} - {None}, set())

    # ── colours ──────────────────────────────────────────────────────────────

    def test_every_value_gets_its_colour(self):
        """Each (sheet, field, word) in the colour table has a rule with that colour's fill and text, on that column."""
        wb = load_workbook(self.build(self.everything()))
        labels = {t: {f: l for l, f in cols} for _, t, cols, _ in excel.SHEETS}
        names = {t: n for n, t, _, _ in excel.SHEETS}
        for (table, field), values in colours.BY_FIELD.items():
            ws = wb[names[table]]
            col = [c.column_letter for c in ws[1] if c.value == labels[table][field]][0]
            for value, colour in values.items():
                with self.subTest(table=table, field=field, value=value):
                    self.assertEqual(look_of(ws, col, value), colours.hexes(colour))

    def test_a_changed_value_gets_the_matching_colour(self):
        """Change a cell's word in the workbook: the rules, read again, give the new word's colour and not the old one's."""
        path = self.build({"actions": [action(1, priority="Urgent")]})
        wb = load_workbook(path)
        ws = wb["Actions"]
        c = self.cell(ws, "DEMO-A-001", "Priority")
        self.assertEqual(look_of(ws, c.column_letter, c.value), colours.hexes("red"))
        c.value = "Low"
        wb.save(path)
        ws = load_workbook(path)["Actions"]
        c = self.cell(ws, "DEMO-A-001", "Priority")
        self.assertEqual(c.value, "Low")
        self.assertEqual(look_of(ws, c.column_letter, c.value), colours.hexes("grey"))
        self.assertEqual(look_of(ws, c.column_letter, "High"), colours.hexes("amber"))

    def test_the_words_stay_and_no_cell_carries_a_static_colour(self):
        """Colour is only ever a rule: the cell keeps its word and has no fill of its own (the sheet reads without colour)."""
        ws = self.sheet({"actions": [action(1, priority="Urgent", due_date="2030-01-01", status="Open")]}, "Actions")
        for label in ("Priority", "Due", "Status"):
            c = self.cell(ws, "DEMO-A-001", label)
            self.assertIn(rgb(c.fill.fgColor), ("000000", None), label)
            self.assertNotIn(rgb(c.font.color), colours.TOKENS.values(), label)
        self.assertEqual(self.cell(ws, "DEMO-A-001", "Priority").value, "Urgent")

    def test_rules_cover_the_data_and_the_spare_rows(self):
        ws = self.sheet({"actions": [action(1), action(2)]}, "Actions")
        ranges = {str(cf.sqref) for cf in ws.conditional_formatting}
        last = 2 + 1 + excel.EXTRA_ROWS
        self.assertIn(f"F2:F{last}", ranges)                  # Priority
        self.assertIn(f"G2:G{last}", ranges)                  # Due (the overdue rule)

    def test_the_colours_themselves(self):
        """RAID follows D-12 (Urgent and High red, Medium amber, Low green); Actions follow the page's .prio pills
        (Urgent red, High amber, Medium and Low grey); the tokens' hex."""
        self.assertEqual({k: colours.PRIORITY[k] for k in ("Urgent", "High", "Medium", "Low")},
                         {"Urgent": "red", "High": "red", "Medium": "amber", "Low": "green"})
        self.assertEqual({k: colours.ACTION_PRIORITY[k] for k in ("Urgent", "High", "Medium", "Low")},
                         {"Urgent": "red", "High": "amber", "Medium": "grey", "Low": "grey"})
        self.assertIs(colours.BY_FIELD[("raid", "priority")], colours.PRIORITY)
        self.assertIs(colours.BY_FIELD[("actions", "priority")], colours.ACTION_PRIORITY)
        self.assertEqual(colours.hexes("red"), ("F6E0DD", "A3301F"))
        self.assertEqual(colours.hexes("amber"), ("FCEFCF", "8A5200"))
        self.assertEqual(colours.hexes("green"), ("D9F1EC", "006B5E"))
        self.assertEqual(colours.MILESTONE_STATE["At risk"], "amber")
        self.assertEqual(colours.DECISION_STATUS, {"Pending": "amber", "Made": "green", "Withdrawn": "grey"})

    def test_action_and_raid_status_follow_the_pages_tones(self):
        self.assertEqual(colours.BY_FIELD[("actions", "status")],
                         {"Open": "blue", "Delayed": "amber", "Closed": "grey", "Cancelled": "grey"})
        self.assertEqual(colours.BY_FIELD[("raid", "status")], {"Open": "blue", "Closed": "grey", "Cancelled": "grey"})
        wb = load_workbook(self.build(self.everything()))
        for sheet, label, table in (("Actions", "Status", "actions"), ("Risks and issues", "Status", "raid")):
            ws = wb[sheet]
            col = [c.column_letter for c in ws[1] if c.value == label][0]
            for value, colour in colours.BY_FIELD[(table, "status")].items():
                self.assertEqual(look_of(ws, col, value), colours.hexes(colour), (sheet, value))

    def test_action_priority_matches_the_pages_prio_pills(self):
        """The Actions mapping is checked against the page's own .prio rules (app.css), so the two cannot drift."""
        import re
        css = (Path(__file__).resolve().parent.parent / "system" / "prototype" / "app.css").read_text(encoding="utf-8")
        family = {"status-red": "red", "status-amber": "amber", "grey-500": "grey", "grey-700": "grey"}
        found = {}
        for level in ("Urgent", "High", "Medium", "Low"):
            m = re.search(r"\.prio\." + level + r"\s*\{\s*background:\s*var\(--([a-z0-9-]+)\)", css)
            self.assertIsNotNone(m, level)
            found[level] = family[m.group(1)]
        self.assertEqual(found, colours.ACTION_PRIORITY)

    def test_every_pair_reads_at_4_5_to_1(self):
        for name in colours.COLOURS:
            fill, text = colours.hexes(name)
            self.assertGreaterEqual(contrast(fill, text), 4.5, name)
        self.assertGreaterEqual(contrast(colours.HEADER_FILL, colours.HEADER_TEXT), 4.5)
        self.assertGreaterEqual(contrast("FFFFFF", colours.LATE_TEXT), 4.5)

    def test_milestone_states(self):
        d = Demo()
        try:
            d.export()
            ws = load_workbook(d.path)["Milestones"]
            for state, ref in d.milestones.items():
                c = self.cell(ws, ref, "Shown state")
                with self.subTest(state=state):
                    self.assertEqual(c.value, state)
                    self.assertEqual(look_of(ws, c.column_letter, state), colours.hexes(colours.MILESTONE_STATE[state]))
            self.assertEqual(set(d.milestones), set(colours.MILESTONE_STATE))        # every state is tested
        finally:
            d.close()

    def test_a_word_the_table_does_not_know_has_no_rule(self):
        ws = self.sheet({"actions": [action(1, priority="Someday")]}, "Actions")
        self.assertIsNone(look_of(ws, "F", "Someday"))
        self.assertIsNone(look_of(ws, "F", ""))

    def test_overdue_is_a_formula_rule_on_the_due_column(self):
        """An open item past its date: a red-text formula rule on Due (Needed by on decisions), true for a real date before
        today on an open row only. The formula is a fixed literal that names columns, never data."""
        wb = load_workbook(self.build(self.everything()))
        want = {"Actions": ("Due", "Status", ("Delayed", "Open")), "Risks and issues": ("Due", "Status", ("Open",)),
                "Decisions": ("Needed by", "Status", ("Pending",))}
        for name, (date_label, status_label, states) in want.items():
            ws = wb[name]
            head = {c.value: c.column_letter for c in ws[1]}
            late = [(col, r) for col, r in rules_of(ws) if r.type == "expression"]
            with self.subTest(sheet=name):
                self.assertEqual(len(late), 1)
                col, rule = late[0]
                self.assertEqual(col, head[date_label])
                d, st = head[date_label], head[status_label]
                opens = ",".join(f'${st}2="{s}"' for s in states)
                self.assertEqual(rule.formula, [f"AND(ISNUMBER(${d}2),${d}2<TODAY(),OR({opens}))"])
                self.assertEqual(rgb(rule.dxf.font.color), colours.LATE_TEXT)
                self.assertIsNone(rule.dxf.fill)                # a red date, not a filled cell

    def test_the_overdue_formula_means_what_the_python_rule_means(self):
        """colours.is_late and the sheet's formula are the same rule: checked by evaluating the formula's parts in Python."""
        today = date(2030, 6, 1)
        for status in ("Open", "Delayed", "Closed", "Cancelled"):
            for due in ("2030-05-31", "2030-06-01", "2030-06-02", "soon", None):
                row = {"status": status, "due_date": due}
                try:
                    is_date = date.fromisoformat(due) < today
                except (TypeError, ValueError):
                    is_date = False
                self.assertEqual(colours.is_late("actions", "due_date", row, today), is_date and status in ("Open", "Delayed"))

    # ── nothing else changes ─────────────────────────────────────────────────

    def test_no_formula_in_any_cell(self):
        """Conditional-format rules hold fixed formulas; no cell does. A value that looks like a formula is only text."""
        rows = self.everything()
        rows["actions"].append(action(2, description="=1+1", priority="Urgent", due_date="2030-01-01", status="Open"))
        path = self.build(rows)
        wb = load_workbook(path)
        for ws in wb:
            for r in ws.iter_rows():
                for c in r:
                    self.assertNotEqual(c.data_type, "f", (ws.title, c.coordinate))
            if ws.title != "Lists":
                self.assertGreater(len(ws.conditional_formatting), 0, ws.title)
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if n.startswith("xl/worksheets/sheet"):
                    xml = z.read(n).decode("utf-8")
                    self.assertNotIn("<f>", xml)
                    self.assertNotIn("<f ", xml)

    def test_updated_is_left_as_the_stored_text(self):
        """Updated is the importer's clash check: it stays the stored UTC text, not a local time."""
        ws = self.sheet({"actions": [action(1)]}, "Actions")
        c = self.cell(ws, "DEMO-A-001", "Updated")
        self.assertEqual((c.value, c.data_type), ("2030-01-01 09:30:00.123", "s"))

    def test_the_import_reads_the_same_values_as_before(self):
        d = Demo()
        try:
            d.export()
            back = excel.read_workbook(d.path)
            for _, table, cols, _ in excel.SHEETS:
                fields = [f for _, f in cols if f not in excel.DERIVED and f != "labels"]
                stored = {r["id"]: {f: r[f] for f in fields} for r in d.store.read(table, "ABC")}
                got = {r["id"]: {f: r.get(f) for f in fields} for r in back[table]}
                self.assertEqual(got, stored, table)
        finally:
            d.close()

    def test_coloured_rows_come_back_as_they_went_in(self):
        rows = {"actions": [action(1, priority="Urgent", due_date="2030-01-01", status="Open"),
                            action(2, priority="Low", status="Closed", closed="2030-01-02")],
                "raid": [raid(1, impact="High", likelihood="Low", priority="High", status="Open")],
                "decisions": [decision(1, status="Pending", needed_by="2030-01-01")],
                "use_cases": [use_case(1, status="High risk")]}
        back = excel.read_workbook(self.build(rows))
        for table, want in rows.items():
            for got, row in zip(back[table], want):
                for field, value in got.items():
                    if field != "labels":
                        self.assertEqual(value, row.get(field), (table, field))

    def test_the_same_rows_give_the_same_bytes_with_all_the_colours(self):
        rows = self.everything()
        rows["actions"] = [action(i, priority=p, due_date="2030-01-01", status="Open")
                           for i, p in enumerate(("Urgent", "High", "Medium", "Low"), start=1)]
        a, b = self.build(rows), self.build(rows)
        self.assertEqual(a.read_bytes(), b.read_bytes())

    def test_values_and_drop_downs_are_the_same_as_without_colour(self):
        """Colour is the only difference: the same export with the colour table emptied has the same values and drop-downs."""
        rows = self.everything()
        rows["actions"] = [action(1, priority="Urgent", due_date="2030-01-01", status="Open")]
        saved = (dict(colours.BY_FIELD), dict(colours.LATE))
        try:
            colours.BY_FIELD.clear()
            colours.LATE.clear()
            plain_book = load_workbook(self.build(rows))
        finally:
            colours.BY_FIELD.update(saved[0])
            colours.LATE.update(saved[1])
        coloured = load_workbook(self.build(rows))
        self.assertEqual(sum(len(ws.conditional_formatting) for ws in plain_book), 0)
        for name in coloured.sheetnames:
            a, b = coloured[name], plain_book[name]
            self.assertEqual([[c.value for c in r] for r in a.iter_rows()], [[c.value for c in r] for r in b.iter_rows()])
            self.assertEqual([(v.type, v.formula1, str(v.sqref)) for v in a.data_validations.dataValidation],
                             [(v.type, v.formula1, str(v.sqref)) for v in b.data_validations.dataValidation])


if __name__ == "__main__":
    unittest.main()
