"""#38: Excel import. A workbook is compared with the database by id; new, changed and removed rows are staged; only the accepted
ones are applied, as one step with one Undo; a row that breaks a rule is reported and nothing is half-applied."""
import shutil
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from _demo import USER, DemoStore
import _workbook as wbk
from clark import excel_import as xi
from clark.store import Invalid, columns
from openpyxl import load_workbook


class ImportCase(unittest.TestCase):
    def setUp(self):
        self.demo = DemoStore()
        self.s = self.demo.store
        self.tmp = Path(tempfile.mkdtemp(prefix="clark-xlsx-"))
        self.a1 = self.demo.action(description="Draft the cut-over plan", due_date="2026-11-01", latest_update="First draft")
        self.a2 = self.demo.action(description="Book the training room")
        self.demo.update("actions", self.a2, {"status": "Closed"})
        self.a3 = self.demo.action(description="Agree the test dates", workstream="Build")
        with self.s.step("dashboard", USER) as st:
            self.r1 = st.create("raid", {"project": "ABC", "type": "Risk", "description": "Key person leaves", "raised": "2026-10-01",
                                         "impact": "High", "likelihood": "Medium", "source": "demo"})
            self.d1 = st.create("decisions", {"project": "ABC", "decision": "Use the shared drive", "date": "2026-10-02",
                                              "status": "Made", "source": "demo"})
        self.path = wbk.make_workbook(self.s, self.tmp / "ABC.xlsx")

    def tearDown(self):
        self.demo.close()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def plan(self):
        return xi.read_workbook(self.s, self.path, "ABC", allow_unrecorded=True)

    def steps(self):
        return self.s.con.execute("SELECT count(*) FROM applies").fetchone()[0]

    def snapshot(self):
        return {t: self.s.read(t) for t in ("actions", "raid", "decisions", "milestones", "use_cases")}

    def kinds(self, plan):
        return [(c.kind, c.id or c.row) for c in plan.changes]


class Staging(ImportCase):
    def test_raised_after_the_closing_date_shows_a_plain_sentence_not_database_text(self):
        """#135: the plan line used to read 'CHECK constraint failed: closed IS NULL OR closed >= raised'."""
        wbk.edit(self.path, "Actions", self.a2, raised=date(2999, 1, 1))
        p = self.plan()
        (c,) = p.changes
        self.assertEqual(c.problem.split(": ", 1)[-1], "This item can't be closed or cancelled before the day it was raised. "
                                                       "Check the raised date, then try again.")
        self.assertNotIn("CHECK constraint", c.problem)

    def test_an_untouched_export_has_no_differences(self):
        p = self.plan()
        self.assertEqual((p.changes, p.problems), ([], []))
        self.assertEqual(p.unchanged, 7)          # 3 actions, 1 risk, 1 decision, 1 milestone, 1 use case

    def test_a_changed_row_is_staged_with_before_and_after_and_nothing_is_written(self):
        before = self.snapshot()
        wbk.edit(self.path, "Actions", self.a1, description="Draft the cut-over plan v2", due_date=date(2026, 12, 4))
        p = self.plan()
        (c,) = p.changes
        self.assertEqual((c.kind, c.id, c.state), ("changed", self.a1, "clean"))
        self.assertEqual(c.after, {"description": "Draft the cut-over plan v2", "due_date": "2026-12-04"})
        self.assertEqual(c.before["due_date"], "2026-11-01")
        self.assertIn("Due 01-11-2026 -> 04-12-2026", c.describe())
        self.assertNotIn("due_date", c.describe())
        self.assertEqual(self.snapshot(), before)

    def test_a_row_with_no_id_is_new(self):
        wbk.add_row(self.path, "Actions", description="Order the laptops", owner="Sam Example", due_date=date(2026, 11, 20))
        (c,) = self.plan().changes
        self.assertEqual((c.kind, c.id, c.table), ("new", "", "actions"))
        self.assertEqual(c.after["project"], "ABC")
        self.assertIn("Excel import", c.after["source"])

    def test_a_row_missing_from_the_workbook_is_staged_as_removed(self):
        wbk.drop_row(self.path, "Actions", self.a3)
        (c,) = self.plan().changes
        self.assertEqual((c.kind, c.id), ("removed", self.a3))
        self.assertIn("No longer in the workbook", c.describe())
        self.assertIn("it is hidden, not deleted. Undo brings it back.", c.describe())

    def test_a_finished_item_missing_from_the_workbook_is_left_alone(self):
        wbk.drop_row(self.path, "Actions", self.a2)           # Closed already
        self.assertEqual(self.plan().changes, [])

    def test_a_change_to_an_archived_row_is_a_problem_at_read_time_in_clean_words(self):
        with self.s.step("dashboard", USER) as st:
            st.archive("actions", self.a1)
        wbk.edit(self.path, "Actions", self.a1, description="Draft the cut-over plan v2")
        p = self.plan()
        self.assertEqual(p.changes, [])
        (problem,) = p.problems
        self.assertEqual(problem.id, self.a1)
        self.assertEqual(problem.message.count(self.a1), 1, "the ID is said once")
        self.assertIn("is archived", problem.message)
        self.assertIn("Restore it in clark first", problem.message)
        self.assertEqual(self.s.read("actions", ref=self.a1)["description"], "Draft the cut-over plan")

    def test_an_unchanged_archived_row_is_not_a_problem(self):
        with self.s.step("dashboard", USER) as st:
            st.archive("actions", self.a1)
        self.assertEqual(self.plan().problems, [])

    def test_an_archived_row_is_not_staged_as_removed_again(self):
        wbk.drop_row(self.path, "Actions", self.a3)
        p = self.plan()
        p.accept(*[c.n for c in p.changes])
        self.assertTrue(xi.apply_import(self.s, p, USER).applied)
        self.assertEqual(self.plan().changes, [])

    def test_the_applied_message_says_in_plain_words_what_was_archived(self):
        wbk.drop_row(self.path, "Actions", self.a3)
        p = self.plan()
        p.accept(*[c.n for c in p.changes])
        res = xi.apply_import(self.s, p, USER)
        self.assertTrue(res.applied, res.message)
        self.assertIn("1 archived because they are no longer in the workbook (nothing was deleted), as one step you can undo.", res.message)
        self.assertNotIn("(archived)", res.message)

    def test_accepting_one_removed_row_archives_only_that_row(self):
        wbk.drop_row(self.path, "Actions", self.a3)
        wbk.drop_row(self.path, "Decisions", self.d1)
        p = self.plan()
        p.accept(*[c.n for c in p.changes if c.id == self.a3])
        res = xi.apply_import(self.s, p, USER)
        self.assertTrue(res.applied, res.message)
        self.assertEqual(self.s.read("actions", ref=self.a3)["archived"], 1)
        self.assertEqual(self.s.read("decisions", ref=self.d1)["archived"], 0)
        self.s.undo(res.step_id, USER)
        self.assertEqual(self.s.read("actions", ref=self.a3)["archived"], 0)

    def test_a_sheet_that_is_missing_removes_nothing(self):
        wbk.make_workbook(self.s, self.path, sheets=("actions",))
        p = self.plan()
        self.assertEqual(p.changes, [])
        self.assertTrue(any("no RAID sheet" in n for n in p.notes))

    def test_a_use_case_missing_from_the_workbook_is_only_a_note(self):
        wbk.drop_row(self.path, "Use cases", self.demo.use_case)
        p = self.plan()
        self.assertEqual(p.changes, [])
        self.assertTrue(any(self.demo.use_case in n and "stays" in n for n in p.notes))

    def test_headers_read_in_any_case_and_with_spaces(self):
        wbk.rename_header(self.path, "Actions", "due_date", "Due date")
        wbk.edit(self.path, "Actions", self.a1, **{"Due date": date(2026, 12, 4)})
        (c,) = self.plan().changes
        self.assertEqual(c.after, {"due_date": "2026-12-04"})

    def test_no_id_column_skips_the_sheet_and_says_so(self):
        wbk.rename_header(self.path, "Actions", "id", "ref")
        p = self.plan()
        self.assertTrue(any("no id column" in x.message for x in p.problems))
        self.assertEqual([c for c in p.changes if c.table == "actions"], [])

    def test_a_file_that_is_not_a_workbook_is_refused_plainly(self):
        bad = self.tmp / "notes.xlsx"
        bad.write_text("not a workbook", encoding="utf-8")
        with self.assertRaises(Invalid) as cm:
            xi.read_workbook(self.s, bad, "ABC")
        self.assertIn("notes.xlsx", str(cm.exception))


    def test_a_workbook_with_only_wrong_sheet_names_has_nothing_to_compare_and_removes_nothing(self):
        wb = load_workbook(str(self.path))
        for ws in wb.worksheets:
            ws.title = "Old " + ws.title
        wb.save(str(self.path))
        p = self.plan()
        self.assertEqual(p.changes, [])
        self.assertEqual(len(p.problems), 1)
        self.assertIn("none of the sheets", p.problems[0].message)

    def test_a_formula_cell_is_left_out_and_noted_not_taken_as_blank(self):
        wbk.edit(self.path, "Actions", self.a1, owner="=1+1")                 # typed as a formula in Excel
        wbk.add_row(self.path, "Actions", description="=A1")                  # a row whose only content is a formula
        p = self.plan()
        self.assertEqual(p.changes, [])                                       # a1's owner is not cleared or changed
        self.assertTrue(any("formula cell was ignored in row" in n and "Actions" in n for n in p.notes), p.notes)

    def test_a_date_with_text_after_it_is_not_a_date(self):
        for good in ("2026-12-04", "2026-12-04 09:30", "2026-12-04T09:30:15", "2026-12-04 09:30:15.25"):
            self.assertEqual(xi.read_date(good), "2026-12-04", good)
        for bad in ("2026-10-28 banana", "2026-12-04 9:30", "2026-12-04 25:00", "2026-12-04T09:30 pm", "2026-12-04 09:30:15.",
                    "2026-12-04 09"):
            with self.assertRaises(ValueError, msg=bad):
                xi.read_date(bad)
        wbk.edit(self.path, "Actions", self.a1, due_date="2026-10-28 banana")
        p = self.plan()
        self.assertEqual(p.changes, [])
        self.assertIn("Due must be a date", p.problems[0].message)

    def test_what_a_person_reads_uses_the_export_labels_and_correct_grammar(self):
        wbk.edit(self.path, "Milestones", self.demo.milestone, baseline_date=date(2030, 6, 1))
        wbk.edit(self.path, "Actions", self.a1, due_date="soon", owner=None)
        wbk.add_row(self.path, "Actions", id="banana7")
        wbk.add_row(self.path, "RAID", id="banana7")
        wbk.add_row(self.path, "Decisions", id="banana7", decision="x")
        wbk.rename_header(self.path, "Use cases", "headline", "Strapline")
        p = self.plan()
        said = " | ".join([x.message for x in p.problems] + p.notes)
        self.assertIn("isn't an action id", said)
        self.assertIn("isn't a RAID item id", said)
        self.assertIn("isn't a decision id", said)
        self.assertIn("(Baseline)", said)
        self.assertIn("Due must be a date", said)
        self.assertIn("Owner is blank", said)
        for raw in ("due_date", "baseline_date", "isn't action id"):
            self.assertNotIn(raw, said)
        wbk.edit(self.path, "Actions", self.a1, due_date=date(2030, 1, 2), owner="Sam Example")
        text = self.plan().report()
        self.assertIn("Owner TBC -> Sam Example", text)
        self.assertNotIn("due_date", text)
        self.assertNotIn("owner ", text)


class SchemaReading(ImportCase):
    def test_the_nullability_read_from_schema_sql_covers_every_column_and_gets_the_known_ones_right(self):
        for t in xi.WORDS:
            self.assertEqual(sorted(xi._table_info(self.s, t)), sorted(columns(self.s.con, t)), t)
        a = xi._table_info(self.s, "actions")
        self.assertEqual((a["due_date"][1], a["owner"][1], a["description"][1]), (False, True, True))
        self.assertEqual((a["client_message"][1], a["client_message"][2]), (True, "''"))
        self.assertEqual(a["client_visible"][0], "INTEGER")
        self.assertFalse(xi._table_info(self.s, "decisions")["date"][1])
        self.assertTrue(xi._is_blank_default(xi._table_info(self.s, "use_cases")["headline"]))


class Reported(ImportCase):
    def test_unknown_id_wrong_pattern_other_project_and_duplicate_are_reported_not_staged(self):
        wbk.add_row(self.path, "Actions", id="ABC-D-001", description="Wrong sheet")
        wbk.add_row(self.path, "Actions", id="banana7", description="Not an id")
        wbk.add_row(self.path, "Actions", id="XYZ-A-001", description="Other project")
        wbk.add_row(self.path, "Actions", id=self.a1, description="Twice")
        p = self.plan()
        self.assertEqual(p.changes, [])
        said = " | ".join(x.message for x in p.problems)
        self.assertEqual(len(p.problems), 4)
        for part in ("ABC-D-001", "banana7", "XYZ-A-001 belongs to project XYZ",
                     f"{self.a1} appears twice"):
            self.assertIn(part, said)
        self.assertEqual({x.field for x in p.problems}, {"id"})

    def test_a_row_naming_another_project_is_reported_with_the_field(self):
        wbk.edit(self.path, "Actions", self.a1, project="XYZ")
        p = self.plan()
        self.assertEqual((p.changes, [(x.id, x.field) for x in p.problems]), ([], [(self.a1, "project")]))

    def test_dates_as_excel_dates_and_as_text_are_both_read(self):
        wbk.edit(self.path, "Actions", self.a1, due_date=datetime(2026, 12, 4, 0, 0))
        wbk.edit(self.path, "Actions", self.a3, due_date="05-12-2026")
        wbk.add_row(self.path, "Actions", description="Third form", due_date="2026-12-06")
        wbk.add_row(self.path, "Actions", description="Slashes", due_date="7/12/2026")
        got = {c.id or c.after["description"]: c.after["due_date"] for c in self.plan().changes}
        self.assertEqual(got, {self.a1: "2026-12-04", self.a3: "2026-12-05", "Third form": "2026-12-06", "Slashes": "2026-12-07"})

    def test_a_date_that_is_not_a_date_is_reported_naming_the_id_and_field(self):
        wbk.edit(self.path, "Actions", self.a1, due_date="31-02-2026")
        wbk.edit(self.path, "Actions", self.a3, due_date="soon")
        p = self.plan()
        self.assertEqual(p.changes, [])
        self.assertEqual(sorted((x.id, x.field) for x in p.problems), sorted([(self.a1, "due_date"), (self.a3, "due_date")]))
        self.assertTrue(all("DD-MM-YYYY" in x.message and x.id in x.message for x in p.problems))

    def test_a_blank_cell_clears_a_value_that_can_be_empty(self):
        wbk.edit(self.path, "Actions", self.a1, due_date=None, latest_update=None)
        (c,) = self.plan().changes
        self.assertEqual(c.after, {"due_date": None, "latest_update": None})

    def test_a_blank_cell_for_a_column_that_must_have_a_value_is_a_problem(self):
        wbk.edit(self.path, "Actions", self.a1, owner=None, description=None)
        p = self.plan()
        self.assertEqual(p.changes, [])
        self.assertEqual(sorted(x.field for x in p.problems), ["description", "owner"])
        self.assertTrue(all(self.a1 in x.message and "can't be empty" in x.message for x in p.problems))

    def test_a_blank_cell_beside_nothing_in_the_database_is_no_change(self):
        wbk.edit(self.path, "Actions", self.a3, due_date=None, client_message=None)     # both already empty
        self.assertEqual(self.plan().changes, [])

    def test_a_column_that_must_hold_text_but_may_be_empty_is_cleared_to_empty_text(self):
        self.demo.update("actions", self.a3, {"client_message": "Dates agreed"})
        wbk.make_workbook(self.s, self.path)
        wbk.edit(self.path, "Actions", self.a3, client_message=None)
        (c,) = self.plan().changes
        self.assertEqual(c.after, {"client_message": ""})

    def test_an_absent_column_is_not_compared(self):
        wbk.make_workbook(self.s, self.path)
        wb = load_workbook(str(self.path))
        ws = wb["Actions"]
        col = [c.column for c in ws[1] if c.value == "latest_update"][0]
        ws.delete_cols(col)
        wb.save(str(self.path))
        self.assertEqual(self.plan().changes, [])

    def test_a_milestones_agreed_date_change_is_ignored_with_a_note(self):
        wbk.edit(self.path, "Milestones", self.demo.milestone, baseline_date=date(2030, 6, 1), owner="Sam Example")
        p = self.plan()
        (c,) = p.changes
        self.assertEqual(c.after, {"owner": "Sam Example"})
        self.assertTrue(any("rebaseline" in n and self.demo.milestone in n for n in p.notes))

    def test_a_clash_is_flagged_and_accept_all_leaves_it_for_the_person(self):
        wbk.edit(self.path, "Actions", self.a1, owner="Sam Example")
        self.demo.update("actions", self.a1, {"priority": "High"})          # changed in clark after the export
        wbk.edit(self.path, "Actions", self.a3, owner="Alex Example")
        p = self.plan()
        by = {c.id: c for c in p.changes}
        self.assertEqual((by[self.a1].state, by[self.a1].clash, by[self.a3].state), ("clash", True, "clean"))
        p.accept_all()
        self.assertEqual((by[self.a1].state, by[self.a3].state), ("clash", "accepted"))
        p.accept(by[self.a1].n)
        self.assertEqual(by[self.a1].state, "accepted")


class Applying(ImportCase):
    def test_only_the_accepted_changes_are_applied(self):
        wbk.edit(self.path, "Actions", self.a1, owner="Sam Example")
        wbk.edit(self.path, "Actions", self.a3, owner="Alex Example")
        wbk.add_row(self.path, "Actions", description="Order the laptops")
        wbk.drop_row(self.path, "RAID", self.r1)
        p = self.plan()
        self.assertEqual(len(p.changes), 4)
        for c in p.changes:
            if c.id == self.a1 or c.kind == "removed":
                p.accept(c.n)
            else:
                p.reject(c.n)
        res = xi.apply_import(self.s, p, USER)
        self.assertTrue(res.applied, res.message)
        self.assertEqual(res.counts, {"new": 0, "changed": 1, "removed": 1})
        self.assertEqual(self.s.read("actions", ref=self.a1)["owner"], "Sam Example")
        self.assertEqual(self.s.read("actions", ref=self.a3)["owner"], "TBC")
        self.assertEqual(len(self.s.read("actions", project="ABC")), 3)       # the new row was rejected
        r = self.s.read("raid", ref=self.r1)
        self.assertEqual((r["status"], r["closed"] is None, r["archived"]), ("Open", True, 1))     # archived, not closed

    def test_nothing_staged_is_applied_until_it_is_accepted(self):
        wbk.edit(self.path, "Actions", self.a1, owner="Sam Example")
        before, n = self.snapshot(), self.steps()
        res = xi.apply_import(self.s, self.plan(), USER)
        self.assertFalse(res.applied)
        self.assertEqual((self.snapshot(), self.steps()), (before, n))
        self.assertIn("Nothing was accepted", res.message)

    def test_new_changed_and_removed_rows_are_one_step_with_one_undo(self):
        wbk.edit(self.path, "Actions", self.a1, owner="Sam Example", due_date=None)
        wbk.edit(self.path, "Decisions", self.d1, decision="Use the team site")
        wbk.add_row(self.path, "Actions", description="Order the laptops", owner="Alex Example")
        wbk.add_row(self.path, "RAID", type="Issue", description="Server is late", impact="Low", likelihood="Low")
        wbk.drop_row(self.path, "Actions", self.a3)
        before, n = self.snapshot(), self.steps()
        p = self.plan()
        p.accept_all()
        self.assertEqual(p.counts(p.accepted())["removed"], 0)                     # accept_all leaves "removed" rows out
        p.accept(*[c.n for c in p.changes if c.kind == "removed"])                 # each is accepted by number
        res = xi.apply_import(self.s, p, USER)
        self.assertTrue(res.applied, res.message)
        self.assertEqual(res.counts, {"new": 2, "changed": 2, "removed": 1})
        self.assertEqual(self.steps(), n + 1)
        self.assertEqual(res.step_id, self.demo.last_step())
        new = [a for a in self.s.read("actions", project="ABC") if a["description"] == "Order the laptops"]
        self.assertEqual((len(new), new[0]["id"], new[0]["source"].startswith("Excel import")), (1, "ABC-A-006", True))
        a3 = self.s.read("actions", ref=self.a3)
        self.assertEqual((a3["status"], a3["archived"]), ("Open", 1))
        self.assertEqual(self.s.con.execute("SELECT DISTINCT route, place FROM item_events WHERE apply_id = ?",
                                            (res.step_id,)).fetchall()[0][:], ("excel", "excel"))
        out = self.s.undo(res.step_id, USER)
        self.assertEqual(out["skipped"], [])
        after = self.snapshot()
        for t in before:
            strip = lambda rows: sorted(({k: v for k, v in r.items() if k != "updated_at"} for r in rows), key=lambda r: r["id"])
            self.assertEqual(strip(after[t]), strip(before[t]), t)

    def test_a_decision_and_a_milestone_missing_from_the_workbook_are_archived(self):
        wbk.drop_row(self.path, "Decisions", self.d1)
        wbk.drop_row(self.path, "Milestones", self.demo.milestone)
        p = self.plan()
        p.accept_all()
        self.assertEqual(p.accepted(), [])
        self.assertIn("2 rows are no longer in the workbook. Accept each one to archive it (nothing is deleted)", p.report())
        p.accept(*[c.n for c in p.changes])
        self.assertTrue(xi.apply_import(self.s, p, USER).applied)
        for t, rid in (("decisions", self.d1), ("milestones", self.demo.milestone)):
            row = self.s.read(t, ref=rid)
            self.assertEqual(row["archived"], 1)
            self.assertNotIn(row["status"], ("Withdrawn", "Cancelled"))

    def test_a_bad_accepted_row_is_reported_and_nothing_is_half_applied(self):
        wbk.edit(self.path, "Actions", self.a1, owner="Sam Example")                 # fine
        wbk.edit(self.path, "Actions", self.a3, status="Finished")                   # breaks a rule
        wbk.edit(self.path, "RAID", self.r1, impact="Enormous")                      # breaks another
        wbk.add_row(self.path, "Actions", description="Order the laptops")           # fine
        p = self.plan()
        bad = {c.id for c in p.changes if c.problem}
        self.assertEqual(bad, {self.a3, self.r1})
        p.accept(*[c.n for c in p.changes])                                          # accept all, bad ones too
        before, n = self.snapshot(), self.steps()
        res = xi.apply_import(self.s, p, USER)
        self.assertFalse(res.applied)
        self.assertEqual((self.snapshot(), self.steps()), (before, n))
        self.assertIsNone(res.step_id)
        self.assertEqual({(x.id, x.field) for x in res.problems}, {(self.a3, "status"), (self.r1, "impact")})
        self.assertIn(self.a3, res.message)
        self.assertIn("Nothing was applied", res.message)

    def test_rejecting_the_bad_row_lets_the_rest_through(self):
        wbk.edit(self.path, "Actions", self.a1, owner="Sam Example")
        wbk.edit(self.path, "Actions", self.a3, status="Finished")
        p = self.plan()
        p.accept_all()                                  # accept_all leaves the bad row waiting
        self.assertEqual({c.id: c.state for c in p.changes}, {self.a1: "accepted", self.a3: "clean"})
        self.assertTrue(xi.apply_import(self.s, p, USER).applied)
        self.assertEqual(self.s.read("actions", ref=self.a3)["status"], "Open")

    def test_a_new_row_that_lacks_what_the_writer_needs_is_reported(self):
        wbk.add_row(self.path, "Actions", owner="Sam Example")                       # no description
        p = self.plan()
        (c,) = p.changes
        self.assertIn("Description", c.problem)
        p.accept(c.n)
        res = xi.apply_import(self.s, p, USER)
        self.assertFalse(res.applied)
        self.assertIn(f"new action (row {c.row} of Actions)", res.message)

    def test_a_change_made_in_clark_after_the_import_was_checked_is_refused(self):
        wbk.edit(self.path, "Actions", self.a1, owner="Sam Example")
        p = self.plan()
        p.accept_all()
        self.demo.update("actions", self.a1, {"priority": "High"})
        res = xi.apply_import(self.s, p, USER)
        self.assertFalse(res.applied)
        self.assertIn("run the import again", res.message)
        self.assertEqual(self.s.read("actions", ref=self.a1)["owner"], "TBC")

    def test_report_reads_in_plain_words(self):
        wbk.edit(self.path, "Actions", self.a1, owner="Sam Example")
        wbk.add_row(self.path, "Actions", id="banana7", description="Ghost")
        text = self.plan().report()
        self.assertIn("0 new, 1 changed, 0 no longer in the workbook, 6 unchanged, 1 reported", text)
        self.assertIn("Not imported:", text)


if __name__ == "__main__":
    unittest.main()
