"""A small workbook maker for the Excel import tests (#38), in the shape the import expects (see src/clark/excel_import.py):
sheets Actions, RAID, Decisions, Milestones, Use cases; row 1 the table's column names as store.columns() gives them; dates
as real Excel dates; flags as Yes or No. It stands in for the export (#37) until that is merged."""
import re
import zipfile
from datetime import date

import _paths  # noqa: F401
from clark.store import DATE_FIELDS, columns
from openpyxl import Workbook, load_workbook

SHEETS = (("Actions", "actions"), ("RAID", "raid"), ("Decisions", "decisions"), ("Milestones", "milestones"),
          ("Use cases", "use_cases"))
FLAGS = {"client_visible", "state_set", "active"}


def cell_of(field, value):
    """A stored value as the export puts it in a cell."""
    if value is None:
        return None
    if field in DATE_FIELDS:
        return date.fromisoformat(value)
    if field in FLAGS:
        return "Yes" if int(value) else "No"
    return value


def make_workbook(store, path, project="ABC", sheets=None):
    """Write every row of `project` to `path`; `sheets` limits which sheets (by table name). Returns the path."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, table in SHEETS:
        if sheets is not None and table not in sheets:
            continue
        ws = wb.create_sheet(name)
        cols = columns(store.con, table)
        ws.append(cols)
        for r in sorted(store.read(table, project=project), key=lambda r: r["id"]):
            ws.append([cell_of(c, r[c]) for c in cols])
    wb.save(str(path))
    return path


def _sheet(wb, name):
    return wb[name]


def _col(ws, field):
    for c in ws[1]:
        if c.value == field:
            return c.column
    raise KeyError(field)


def _row_of(ws, item_id):
    col = _col(ws, "id")
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, col).value == item_id:
            return r
    raise KeyError(item_id)


def edit(path, sheet, item_id, **cells):
    """Set cells of the row with this id (None clears a cell)."""
    wb = load_workbook(str(path))
    ws = _sheet(wb, sheet)
    r = _row_of(ws, item_id)
    for f, v in cells.items():
        ws.cell(r, _col(ws, f)).value = v
    wb.save(str(path))


def add_row(path, sheet, **cells):
    """Append a row (leave `id` out for a new item)."""
    wb = load_workbook(str(path))
    ws = _sheet(wb, sheet)
    ws.append([cells.get(c.value) for c in ws[1]])
    wb.save(str(path))


def drop_row(path, sheet, item_id):
    wb = load_workbook(str(path))
    ws = _sheet(wb, sheet)
    ws.delete_rows(_row_of(ws, item_id))
    wb.save(str(path))


def rename_header(path, sheet, old, new):
    wb = load_workbook(str(path))
    ws = _sheet(wb, sheet)
    ws.cell(1, _col(ws, old)).value = new
    wb.save(str(path))


CUSTOM = "docProps/custom.xml"
_CUSTOM_TYPE = "application/vnd.openxmlformats-officedocument.custom-properties+xml"
_CUSTOM_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties"


def set_custom_part(path, data):
    """Rewrite the workbook at `path` with `data` (bytes) as its custom-properties part, or with none when `data` is None (what a
    program that drops custom properties writes). The content types and the package relationships follow."""
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist() if n != CUSTOM}
    types = re.sub(rb'<Override PartName="/docProps/custom.xml"[^>]*/>', b"", parts["[Content_Types].xml"])
    rels = re.sub(rb'<Relationship [^>]*Target="docProps/custom.xml"[^>]*/>', b"", parts["_rels/.rels"])
    if data is not None:
        types = types.replace(b"</Types>", f'<Override PartName="/{CUSTOM}" ContentType="{_CUSTOM_TYPE}"/></Types>'.encode())
        rels = rels.replace(b"</Relationships>", f'<Relationship Type="{_CUSTOM_REL}" Target="{CUSTOM}" Id="rIdC"/></Relationships>'.encode())
        parts[CUSTOM] = data
    parts["[Content_Types].xml"], parts["_rels/.rels"] = types, rels
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)


def custom_part(value, name="clark_export_id", kind="lpwstr", repeat=1):
    """The bytes of a custom-properties part holding `repeat` properties `name` of type `kind` with the text `value`."""
    one = (f'<property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="{{pid}}" name="{name}"><vt:{kind}>{value}</vt:{kind}>'
           "</property>")
    body = "".join(one.replace("{pid}", str(2 + i)) for i in range(repeat))
    return ('<?xml version="1.0" encoding="UTF-8"?><Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/'
            'custom-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">' + body
            + "</Properties>").encode("utf-8")
