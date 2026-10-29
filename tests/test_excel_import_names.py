"""#143 (FR-70): the import reads the OLD sheet and header names (RAID, Client messaging, Baseline) and the NEW ones (Risks and
issues, Client wording, Agreed), alone or mixed, so a workbook exported before the rename still imports. Two names for one thing
in one workbook are refused in plain words, and nothing is written."""
import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path

from _demo import USER, DemoStore
from clark import excel
from clark import excel_import as xi
from clark.store import Invalid
from openpyxl import load_workbook

ON = "2030-06-01"


class NameCase(unittest.TestCase):
    def setUp(self):
        self.demo = DemoStore()
        self.s = self.demo.store
        self.tmp = Path(tempfile.mkdtemp(prefix="clark-names-"))
        self.a1 = self.demo.action(description="Draft the cut-over plan", client_message="Plan is coming")
        with self.s.step("dashboard", USER) as st:
            self.r1 = st.create("raid", {"project": "ABC", "type": "Risk", "description": "Key person leaves", "raised": "2026-10-01",
                                         "client_message": "We are watching this", "source": "demo"})
            self.d1 = st.create("decisions", {"project": "ABC", "decision": "Choose the team", "status": "Pending",
                                              "client_message": "Soon", "source": "demo"})
        self.path = self.tmp / "ABC.xlsx"
        self.state = self.tmp / "state.json"
        excel.export_project(self.s, "ABC", self.path, self.state, on=ON)

    def tearDown(self):
        self.demo.close()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def read(self):
        return xi.read_workbook(self.s, self.path, "ABC", state_path=self.state)

    def rename(self, raid_sheet=None, **headers):
        """Give the RAID sheet a name, and headers new names, as `{"Actions": ("Client messaging", "Client wording")}` ."""
        wb = load_workbook(str(self.path))
        if raid_sheet:
            wb[excel.SHEET_OF["raid"]].title = raid_sheet
        for sheet, (old, new) in headers.items():
            ws = wb[sheet]
            for c in ws[1]:
                if c.value == old:
                    c.value = new
        wb.save(str(self.path))

    def edit(self, sheet, item, label, value):
        wb = load_workbook(str(self.path))
        ws = next(w for w in wb.worksheets
                  if w.title.lower().replace("_", " ") in (sheet.lower(), "risks and issues") and (sheet != "RAID" or w.title != "Actions"))
        head = {c.value: c.column for c in ws[1]}
        row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == item)
        ws.cell(row, head[label]).value = value
        wb.save(str(self.path))

    def snapshot(self):
        return {t: self.s.read(t) for t in ("actions", "raid", "decisions", "milestones")}


class Names(NameCase):
    def edited(self, wording, raid="RAID", agreed="Baseline"):
        """Edit the client wording of an action and a risk, and the milestone's agreed date; return the read plan."""
        self.edit("Actions", self.a1, wording, "Plan is on its way")
        self.edit(raid, self.r1, wording, "Still watching")
        self.edit("Milestones", self.demo.milestone, agreed, date(2031, 1, 1))
        return self.read()

    def check(self, p):
        self.assertEqual(p.problems, [])
        got = {(c.table, c.id): c.after for c in p.changes}
        self.assertEqual(got[("actions", self.a1)], {"client_message": "Plan is on its way"})
        self.assertEqual(got[("raid", self.r1)], {"client_message": "Still watching"})
        self.assertEqual([c for c in got if c[0] == "milestones"], [])                 # an agreed date is not changed by an import...
        self.assertTrue([n for n in p.notes if "agreed date" in n], p.notes)           # ...and the note shows the column was read
        p.accept_all()
        self.assertTrue(xi.apply_import(self.s, p, USER).applied)
        self.assertEqual(self.s.read("raid", ref=self.r1)["client_message"], "Still watching")

    def test_a_workbook_with_the_old_names_imports(self):
        self.rename(**{"Actions": ("Client messaging", "Client messaging")})            # as exported today
        self.check(self.edited("Client messaging"))

    def test_a_workbook_with_the_new_names_imports(self):
        self.rename("Risks and issues", Actions=("Client messaging", "Client wording"), Decisions=("Client messaging", "Client wording"),
                    Milestones=("Baseline", "Agreed"))
        wb = load_workbook(str(self.path))
        ws = wb["Risks and issues"]
        next(c for c in ws[1] if c.value == "Client messaging").value = "Client wording"
        wb.save(str(self.path))
        self.check(self.edited("Client wording", raid="Risks and issues", agreed="Agreed"))

    def test_a_mixed_workbook_imports(self):
        self.rename("Risks and issues", Actions=("Client messaging", "Client wording"))
        self.check(self.edited_mixed())

    def edited_mixed(self):
        self.edit("Actions", self.a1, "Client wording", "Plan is on its way")
        self.edit("RAID", self.r1, "Client messaging", "Still watching")
        self.edit("Milestones", self.demo.milestone, "Baseline", date(2031, 1, 1))
        return self.read()

    def test_sheet_names_match_in_any_case_with_spaces_or_underscores(self):
        self.rename("risks_AND_issues")
        self.edit("RAID", self.r1, "Client messaging", "Still watching")
        p = self.read()
        self.assertEqual([(c.table, c.id) for c in p.changes], [("raid", self.r1)])

    def test_an_unchanged_export_in_either_naming_gives_no_changes(self):
        p = self.read()
        self.assertEqual((p.changes, p.problems), ([], []))
        self.rename("Risks and issues", Actions=("Client messaging", "Client wording"), Milestones=("Baseline", "Agreed"))
        p = self.read()
        self.assertEqual((p.changes, p.problems), ([], []))

    def test_an_unknown_sheet_and_header_are_named_and_nothing_partial_is_written(self):
        wb = load_workbook(str(self.path))
        wb.create_sheet("Risks & issues")                          # a look-alike, not a name clark knows
        wb.save(str(self.path))
        self.rename(Actions=("Owner", "Who"))
        before = self.snapshot()
        p = self.read()
        self.assertEqual(p.problems, [])
        self.assertTrue([n for n in p.notes if "Who" in n and "ignored" in n], p.notes)
        self.assertEqual(self.snapshot(), before)


class Duplicates(NameCase):
    def test_the_old_and_new_sheet_name_together_are_refused(self):
        wb = load_workbook(str(self.path))
        wb.copy_worksheet(wb[excel.SHEET_OF["raid"]]).title = "Risks and issues"
        wb.save(str(self.path))
        before = self.snapshot()
        with self.assertRaises(Invalid) as cm:
            self.read()
        said = str(cm.exception)
        self.assertIn("RAID", said)
        self.assertIn("Risks and issues", said)
        self.assertIn("Nothing was changed", said)
        self.assertEqual(self.snapshot(), before)

    def test_the_old_and_new_header_together_are_refused(self):
        wb = load_workbook(str(self.path))
        ws = wb["Milestones"]
        ws.cell(1, ws.max_column + 1).value = "Agreed"             # beside "Baseline"
        wb.save(str(self.path))
        before = self.snapshot()
        with self.assertRaises(Invalid) as cm:
            self.read()
        said = str(cm.exception)
        for part in ("Milestones", "Baseline", "Agreed", "Nothing was changed"):
            self.assertIn(part, said)
        self.assertEqual(self.snapshot(), before)

    def test_the_old_and_new_wording_header_together_are_refused(self):
        wb = load_workbook(str(self.path))
        ws = wb["Actions"]
        ws.cell(1, ws.max_column + 1).value = "Client wording"
        wb.save(str(self.path))
        with self.assertRaises(Invalid):
            self.read()


class Huge(NameCase):
    def test_a_far_cell_on_a_new_named_sheet_is_still_refused(self):
        self.rename("Risks and issues")
        wb = load_workbook(str(self.path))
        wb["Risks and issues"]["XFD1000"] = "x"
        wb.save(str(self.path))
        with self.assertRaises(Invalid):
            self.read()


if __name__ == "__main__":
    unittest.main()
