"""#85: text that starts like a formula stays text through an export and an import, and is never a formula. Fictional data only.

This file holds only what the export does today: the same bytes for the same rows and day, the file's hash in the record, the
look-alike characters and the font and wrap of a text cell left alone, a file from the first release still reading back, and the
export-then-import round trip of text that starts with = + - @ or a tab. It reads the workbook back with openpyxl. It cannot
open Excel, Numbers or LibreOffice.

Run: python3 dev/run_tests.py test_excel_text_cells
"""
import hashlib
import shutil
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock

import _paths  # noqa: F401  (puts src/ and vendor/ on the path)
from _demo import USER, DemoStore
from openpyxl import load_workbook

from clark import excel
from clark import excel_import as xi
from clark import page

FIXED_ID = "0123456789ab4def8123456789abcdef"
ON = "2030-06-01"

_NO_BRAND = tempfile.mkdtemp(prefix="clark-no-brand-")    # a brand logo on this computer never changes a test


def setUpModule():
    global _OLD_BRANDING
    _OLD_BRANDING, page.BRANDING = page.BRANDING, Path(_NO_BRAND)


def tearDownModule():
    page.BRANDING = _OLD_BRANDING
    shutil.rmtree(_NO_BRAND, ignore_errors=True)


# text that could be taken for a formula
GUARDED = ("=SUM(1)", "+1", "-note", "@SUM(A1)", "\tTabbed", "=", "-", "+", "@", "=1+1", '=HYPERLINK("http://x.invalid","hi")',
           "- a bullet", "-- two dashes", "=cmd|' /C calc'!A0")


def action(text, **more):
    return {"id": "DEMO-A-001", "project": "DEMO", "raised": "2030-01-01", "client_message": "", "client_visible": 1,
            "source": "demo", "description": text, **more}


class DataSheets(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp(prefix="clark-text-cells-"))
        self.addCleanup(shutil.rmtree, self.dir, True)
        self.path, self.state = self.dir / "DEMO.xlsx", self.dir / "state.json"

    def export(self, text, **more):
        self.path.unlink() if self.path.exists() else None
        excel.export(self.path, "DEMO", {"actions": [action(text, **more)]}, self.state, on=ON)

    def test_full_width_look_alikes_are_left_alone(self):
        self.export("＝1+1")
        c = load_workbook(self.path)["Actions"]["B2"]
        self.assertEqual((c.value, c.data_type, c.quotePrefix), ("＝1+1", "s", False))

    def test_the_style_keeps_the_font_and_the_wrap(self):
        self.export("=SUM(1)")
        plain = load_workbook(self.path)["Actions"]["B2"]
        self.export("SUM(1)")
        other = load_workbook(self.path)["Actions"]["B2"]
        self.assertEqual((plain.font.name, plain.alignment.wrap_text, plain.alignment.vertical),
                         (other.font.name, other.alignment.wrap_text, other.alignment.vertical))

    def test_the_same_rows_and_day_give_the_same_bytes_with_guarded_text(self):
        rows = {"actions": [action(text) for text in GUARDED[:4]]}
        for i, row in enumerate(rows["actions"], start=1):
            row["id"] = f"DEMO-A-{i:03d}"
        out = []
        for real in (datetime(2031, 1, 1, 9, 0, 1), datetime(2032, 5, 6, 7, 8, 9)):
            with mock.patch.object(excel, "clock", lambda real=real: real), \
                    mock.patch.object(excel, "new_export_id", lambda: FIXED_ID), \
                    mock.patch("time.time", lambda real=real: real.timestamp()), \
                    mock.patch("time.localtime", lambda t=None, real=real: real.timetuple()):
                path = self.dir / f"{len(out)}.xlsx"
                excel.export(path, "DEMO", rows, self.dir / f"state{len(out)}.json", on=ON)
                out.append(path.read_bytes())
        self.assertEqual(out[0], out[1])

    def test_the_record_holds_the_hash_of_the_file_written(self):
        self.export("=SUM(1)")
        record = excel.load_state(self.state, self.path, "DEMO")
        self.assertEqual(record["hash"], hashlib.sha256(self.path.read_bytes()).hexdigest())
        self.assertEqual(record["rows"]["actions"]["DEMO-A-001"]["description"], "=SUM(1)")     # the record holds the text as typed


class OneCell(unittest.TestCase):
    def test_a_file_from_the_first_release_still_reads_back(self):
        # that release wrote =1+1 as '=1+1; the import still takes the apostrophe off
        for old, text in (("'=1+1", "=1+1"), ("'-note", "-note"), ("'@x", "@x"), ("''=1", "'=1"), ("'quoted", "'quoted")):
            self.assertEqual(excel.from_cell("description", old), text)
            self.assertEqual(xi.from_cell("description", old, "TEXT")[0], text)


class RoundTrip(unittest.TestCase):
    """Export, then import: the text is the same, no change is found, and a real formula typed by a person is still refused."""
    TEXTS = GUARDED + ("'=1", "''x", " =1", "＝1+1", "=" + "z" * 9000)

    def setUp(self):
        self.demo = DemoStore()
        self.addCleanup(self.demo.close)
        self.dir = Path(tempfile.mkdtemp(prefix="clark-text-cells-"))
        self.addCleanup(shutil.rmtree, self.dir, True)
        with self.demo.store.step("dashboard", USER) as st:
            self.ids = [st.create("actions", {"project": "ABC", "description": text, "raised": "2026-10-01", "source": "demo"})
                        for text in self.TEXTS]
        self.path, self.state = self.dir / "ABC work items.xlsx", self.dir / "state.json"
        excel.export_project(self.demo.store, "ABC", self.path, self.state, on=ON)

    def plan(self):
        return xi.read_workbook(self.demo.store, self.path, "ABC", state_path=self.state)

    def test_an_unchanged_export_imports_as_no_change(self):
        plan = self.plan()
        self.assertEqual([c for c in plan.changes if c.kind != "unchanged"], [])
        self.assertEqual(plan.problems, [])

    def test_the_text_in_the_file_is_the_text_in_the_store(self):
        ws = load_workbook(self.path)["Actions"]
        head = {c.value: c.column for c in ws[1]}
        by_id = {ws.cell(row=r, column=head["ID"]).value: ws.cell(row=r, column=head["Description"]) for r in range(2, ws.max_row + 1)}
        for item_id, text in zip(self.ids, self.TEXTS):
            self.assertEqual(by_id[item_id].value, excel.to_cell("description", text))
            self.assertEqual(by_id[item_id].data_type, "s")

    def test_a_real_formula_typed_by_a_person_is_refused(self):
        wb = load_workbook(self.path)
        ws = wb["Actions"]
        head = {c.value: c.column for c in ws[1]}
        plain = ws.cell(row=2, column=head["Owner"])
        guarded = ws.cell(row=3, column=head["Owner"])           # next to a cell that has the quote prefix on its text
        plain.value = "=1+1"
        guarded.value = "=2+2"
        self.assertEqual((plain.data_type, guarded.data_type), ("f", "f"))
        wb.save(self.path)
        plan = self.plan()
        self.assertTrue(any("formula cell was ignored" in n for n in plan.notes), plan.notes)
        self.assertEqual([c for c in plan.changes if c.kind != "unchanged"], [], "a formula is never read as a value")

    def test_an_ordinary_edit_next_to_guarded_text_imports_with_the_text_intact(self):
        wb = load_workbook(self.path)
        ws = wb["Actions"]
        head = {c.value: c.column for c in ws[1]}
        ws.cell(row=2, column=head["Owner"]).value = "Sam Example"
        wb.save(self.path)
        plan = self.plan()
        changed = [c for c in plan.changes if c.kind != "unchanged"]
        self.assertEqual(len(changed), 1)
        plan.accept_all()
        xi.apply_import(self.demo.store, plan, USER, state_path=self.state)
        got = {r["id"]: r["description"] for r in self.demo.store.read("actions", "ABC") if r["id"] in self.ids}
        self.assertEqual([got[i] for i in self.ids], list(self.TEXTS))


if __name__ == "__main__":
    unittest.main()
