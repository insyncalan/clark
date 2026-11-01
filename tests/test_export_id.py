"""The export id inside an exported workbook (src/clark/excel.py: new_export_id, export_id_of, id_key, load_state): the import
finds the export's record by it after the workbook was edited or saved again, and a workbook with no id, or a forged, odd or
other project's id, is unrecorded, so every difference is staged for the person to check (it fails safe).

Run: python3 dev/run_tests.py test_export_id
"""
import json
import shutil
import tempfile
import unittest
import uuid
import zipfile
from datetime import date
from pathlib import Path
from unittest import mock

from _paths import SRC  # noqa: F401
from _workbook import custom_part, set_custom_part
from openpyxl import load_workbook
from openpyxl.packaging.custom import IntProperty
from clark import excel, excel_import
import test_publish_route as base
from test_preview_export_route import get

ROWS = {"actions": [{"id": "DEMO-A-001", "project": "DEMO", "description": "Book the room", "status": "Open",
                     "raised": "2031-01-01", "source": "test", "updated_at": "2031-01-01 09:00:00"}]}


def members(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


class TheId(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp(prefix="clark-export-id-"))
        self.addCleanup(shutil.rmtree, self.dir, True)

    def export(self, name, rows=ROWS, project="DEMO"):
        path = self.dir / f"{name}.xlsx"
        excel.export(path, project, rows, self.dir / f"{name}.json", on="2031-01-01")
        return path

    def test_each_export_carries_a_random_version_4_uuid_in_its_custom_properties_only(self):
        a, b = self.export("a"), self.export("b")                  # the same rows, the same day, two records
        ids = [excel.export_id_of(p) for p in (a, b)]
        self.assertNotEqual(ids[0], ids[1], "the id is random, not worked out from the rows")
        for token in ids:
            self.assertEqual(uuid.UUID(hex=token).version, 4)
            self.assertEqual(uuid.UUID(hex=token).hex, token)
        ma, mb = members(a), members(b)
        self.assertEqual(sorted(n for n in ma if ma[n] != mb.get(n)), [excel.CUSTOM_PART], "the id is the only difference")
        for name, data in ma.items():
            if name != excel.CUSTOM_PART:
                self.assertNotIn(ids[0].encode(), data, name)
        wb = load_workbook(a)
        cells = [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]
        self.assertFalse([c for c in cells if ids[0] in c], "never in a cell or on the Cover")

    def test_the_record_keeps_the_id_and_nothing_personal_rides_with_it(self):
        path = self.export("a")
        record = json.loads((self.dir / "a.json").read_text(encoding="utf-8"))["workbooks"][excel._key(path)]
        self.assertEqual(record["export_id"], excel.export_id_of(path))
        part = members(path)[excel.CUSTOM_PART].decode("utf-8")
        self.assertEqual(part.count("<property "), 1, "one property: the id, and no name, path or project")
        self.assertNotIn("DEMO", part)

    def test_an_export_to_the_same_path_with_the_same_rows_keeps_its_id_and_its_bytes(self):
        path, state = self.dir / "w.xlsx", self.dir / "w.json"
        excel.export(path, "DEMO", ROWS, state, on="2031-01-01")
        first = path.read_bytes()
        excel.export(path, "DEMO", ROWS, state, on="2031-01-01")
        self.assertEqual(path.read_bytes(), first)
        changed = {"actions": [dict(ROWS["actions"][0], description="Book a bigger room")]}
        token = excel.export_id_of(path)
        excel.export(path, "DEMO", changed, state, on="2031-01-01")
        self.assertNotEqual(excel.export_id_of(path), token, "another file never shares an id")

    def test_odd_values_are_never_taken_as_an_id(self):
        good = uuid.uuid4().hex
        odd = {"too long": "a" * 100000, "not hex": "z" * 32, "upper case": good.upper(), "with dashes": str(uuid.UUID(hex=good)),
               "markup": "&lt;b&gt;" + good + "&lt;/b&gt;", "a path": "../../" + good, "a NUL": good + "&#0;",
               "not version 4": good[:12] + "1" + good[13:], "too short": good[:31], "not ASCII": good[:31] + "é",
               "space": good + " "}
        base_path = self.export("base")
        for why, value in odd.items():
            with self.subTest(why):
                path = self.dir / "odd.xlsx"
                shutil.copy(base_path, path)
                set_custom_part(path, custom_part(value))
                self.assertIsNone(excel.export_id_of(path))
        for why, part in (("a number", custom_part("7", kind="i4")), ("two of them", custom_part(good, repeat=2)),
                          ("another name", custom_part(good, name="clark_export_id2")), ("not XML", b"<Properties><"),
                          ("raw NUL", custom_part(good).replace(good.encode(), good.encode() + b"\x00")),
                          ("a huge part", custom_part(good) + b" " * (excel.MAX_CUSTOM_BYTES + 1))):
            with self.subTest(why):
                path = self.dir / "odd.xlsx"
                shutil.copy(base_path, path)
                set_custom_part(path, part)
                self.assertIsNone(excel.export_id_of(path))
        self.assertIsNone(excel.export_id_of(self.dir / "missing.xlsx"))
        (self.dir / "text.xlsx").write_text("not a zip", encoding="utf-8")
        self.assertIsNone(excel.export_id_of(self.dir / "text.xlsx"))
        set_custom_part(path, custom_part(good))
        self.assertEqual(excel.export_id_of(path), good, "the helper writes a part the reader does read")


class Case(unittest.TestCase):
    def setUp(self):
        self.rig = base.Rig(self)
        self.state = self.rig.data_dir / excel.STATE_FILE

    def download(self, name="PUB.xlsx", project="PUB"):
        _, _, raw = get(self.rig, f"/api/export?project={project}")
        mine = Path(self.rig.root) / "Downloads" / name
        mine.parent.mkdir(exist_ok=True)
        mine.write_bytes(raw)
        return mine

    def read(self, path, project="PUB"):
        return excel_import.read_workbook(self.rig.store, path, project, state_path=self.state)

    def books(self):
        return json.loads(self.state.read_text(encoding="utf-8"))["workbooks"]


class Found(Case):
    def test_the_record_is_found_with_no_edit_a_resave_one_cell_edited_and_a_column_sorted(self):
        def nothing(wb):
            pass

        def one_cell(wb):
            ws = wb["Milestones"]
            ws.cell(row=2, column=[c.value for c in ws[1]].index("Milestone") + 1).value = "Renamed in Excel"

        def sorted_column(wb):
            ws = max(wb.worksheets[1:], key=lambda w: w.max_row if w.title in ("Actions", "Milestones", "RAID") else 0)
            rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2, max_row=ws.max_row) if any(c.value for c in r)]
            self.assertGreater(len(rows), 1, "a sheet with rows to sort")
            for i, values in enumerate(sorted(rows, key=lambda r: str(r[1]), reverse=True), start=2):
                for j, v in enumerate(values, start=1):
                    ws.cell(row=i, column=j).value = v

        for why, change, staged in (("no edit", None, 0), ("saved again", nothing, 0), ("one cell", one_cell, 1),
                                    ("sorted", sorted_column, 0)):
            with self.subTest(why):
                mine = self.download(f"{why}.xlsx")
                if change is not None:
                    wb = load_workbook(mine)
                    change(wb)
                    wb.save(mine)
                plan = self.read(mine)
                self.assertFalse(plan.unrecorded, why)
                self.assertEqual(len([c for c in plan.changes if c.kind == "changed"]), staged, why)
                self.assertEqual([c for c in plan.changes if c.kind != "changed"], [], why)

    def test_the_route_keeps_the_record_under_the_id_and_the_project_only(self):
        mine = self.download()
        token = excel.export_id_of(mine)
        self.assertEqual([k for k in self.books() if k.startswith(excel.KEYED)], [excel.id_key("PUB", token)])
        for name, data in members(mine).items():
            if name != excel.CUSTOM_PART:
                self.assertNotIn(token.encode(), data, name)

    def test_an_import_writes_back_to_the_id_and_a_short_record_under_the_path(self):
        mine = self.download()
        wb = load_workbook(mine)
        ws = wb["Milestones"]
        ws.cell(row=2, column=[c.value for c in ws[1]].index("Milestone") + 1).value = "Renamed in Excel"
        wb.save(mine)
        plan = self.read(mine)
        plan.accept(*[c.n for c in plan.changes])
        self.assertTrue(excel_import.apply_import(self.rig.store, plan, base.USER, state_path=self.state).applied)
        books = self.books()
        record = books[excel.id_key("PUB", excel.export_id_of(mine))]
        self.assertEqual(record["imported_hash"], excel.file_hash(mine))
        short = books[str(mine.resolve())]
        self.assertEqual(short["imported_hash"], excel.file_hash(mine))
        self.assertNotIn("rows", short, "the short record holds no copy of the rows")
        self.assertNotIn("hash", short)
        self.assertTrue(self.read(mine).already_imported)


class NewRows(Case):
    def add_action(self, path, words):
        wb = load_workbook(path)
        ws = wb["Actions"]
        head = [c.value for c in ws[1]]
        last = max(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value) if ws.max_row > 1 else 1
        ws.cell(row=last + 1, column=head.index("Description") + 1).value = words
        ws.cell(row=last + 1, column=head.index("Project") + 1).value = "PUB"
        ws.cell(row=last + 1, column=head.index("Owner") + 1).value = "Owner A"
        ws.cell(row=last + 1, column=head.index("Status") + 1).value = "Open"
        ws.cell(row=last + 1, column=head.index("Due") + 1 if "Due" in head else head.index("Due date") + 1).value = date(2026, 11, 30)
        wb.save(path)

    def import_all(self, path, confirm=False):
        plan = self.read(path)
        if confirm:
            plan.confirm_unrecorded()
        plan.accept(*[c.n for c in plan.changes])
        got = excel_import.apply_import(self.rig.store, plan, base.USER, state_path=self.state)
        self.assertTrue(got.applied, got.message)
        return plan

    def test_a_new_row_is_listed_once_and_not_made_again_after_the_id_record_is_pruned(self):
        mine = self.download()
        self.add_action(mine, "Added in Excel")
        self.import_all(mine)
        made = excel.created_for(self.state, mine, "PUB")
        self.assertEqual(len(made), 1, "recorded under the id and the path, listed once")
        excel.prune_hash_records(self.state, "2027-06-01")             # long after: the id record is gone
        self.assertEqual([k for k in self.books() if k.startswith(excel.KEYED)], [])
        wb = load_workbook(mine)                                      # edited again, so it is not the file last imported
        ws = wb["Milestones"]
        ws.cell(row=2, column=[c.value for c in ws[1]].index("Milestone") + 1).value = "Renamed in Excel"
        wb.save(mine)
        plan = self.read(mine)
        self.assertTrue(plan.unrecorded, "the export's record went, so it fails safe")
        self.assertEqual(plan.already_created, 1, "the row the import made is not made again")
        self.assertEqual([c for c in plan.changes if c.kind == "new"], [])


class NotAnotherProjects(Case):
    def setUp(self):
        super().setUp()
        with self.rig.store.step("dashboard", base.USER) as st:
            st.create("projects", {"code": "OTH", "name": "Other Programme", "health": "On track"})

    def test_a_workbook_of_one_project_is_no_record_for_another(self):
        mine = self.download()
        self.assertFalse(self.read(mine).unrecorded)
        self.assertTrue(self.read(mine, "OTH").unrecorded)
        self.assertEqual(excel.load_state(self.state, mine, "OTH"), {})

    def test_another_projects_id_copied_into_a_workbook_is_no_record(self):
        theirs = self.download("OTH.xlsx", "OTH")
        mine = self.download()
        set_custom_part(mine, custom_part(excel.export_id_of(theirs)))
        self.assertEqual(excel.export_id_of(mine), excel.export_id_of(theirs))
        plan = self.read(mine)
        self.assertTrue(plan.unrecorded)
        self.assertEqual(excel.load_state(self.state, mine, "PUB"), {})
        self.assertFalse(excel_import.apply_import(self.rig.store, plan, base.USER, state_path=self.state).applied)

    def test_a_forged_id_is_no_record(self):
        mine = self.download()
        set_custom_part(mine, custom_part(uuid.uuid4().hex))
        self.assertTrue(self.read(mine).unrecorded)

    def test_odd_ids_in_a_downloaded_workbook_never_break_the_import(self):
        for value in ("a" * 100000, "&lt;x/&gt;", "../" * 20, uuid.uuid4().hex.upper()):
            with self.subTest(value[:10]):
                mine = self.download()
                set_custom_part(mine, custom_part(value))
                self.assertTrue(self.read(mine).unrecorded)
        mine = self.download()                                        # a NUL is not XML: the file is refused in words, nothing else
        set_custom_part(mine, custom_part("&#0;"))
        with self.assertRaises(excel_import.Invalid):
            self.read(mine)
        mine = self.download()
        wb = load_workbook(mine)
        wb.custom_doc_props = type(wb.custom_doc_props)()
        wb.custom_doc_props.append(IntProperty(name=excel.EXPORT_ID_NAME, value=7))
        wb.save(mine)
        self.assertTrue(self.read(mine).unrecorded)

    def test_a_state_file_of_the_wrong_shape_is_no_record_not_a_crash(self):
        mine = self.download()
        for text in ("[]", '{"workbooks": []}', '{"workbooks": {"x": 1}}'):
            with self.subTest(text):
                self.state.write_text(text, encoding="utf-8")
                self.assertTrue(self.read(mine).unrecorded)


if __name__ == "__main__":
    unittest.main()
