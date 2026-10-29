"""#85b, the pure part: the lines of the workbook's Cover sheet (clark.excel_cover). Fictional data only.

Run: python3 dev/run_tests.py test_excel_cover
"""
import os
import shutil
import sqlite3
import tempfile
import time
import unittest
from pathlib import Path

import _paths  # noqa: F401  (puts src/ and vendor/ on the path)
from _demo import USER, DemoStore
from openpyxl import Workbook, load_workbook

from clark import excel
from clark import excel_colours as colours
from clark import excel_cover as cover
from clark import excel_import as xi

PROJECT = {"code": "DEMO", "name": "Harbour Lights upgrade", "current_state": "<p>Phase two is under way.</p><ul><li>Cabling done</li><li>Testing next</li></ul>",
           "internal_notes": "CANARY-INTERNAL-NOTE", "health_reason": "CANARY-PROJECT-FIELD"}
SECTIONS = [
    {"id": "s1", "type": "current", "title": "Current state", "body": "", "hidden": 0},
    {"id": "s2", "type": "glance", "title": "At a glance", "body": "", "hidden": 0},
    {"id": "s3", "type": "text", "title": "Summary", "body": "<p>On course for the spring opening.</p>", "hidden": 0},
    {"id": "s4", "type": "text", "title": "This week", "body": "<p>Lights tested on the east pier.</p>", "hidden": 0},
    {"id": "s5", "type": "text", "title": "Hidden section", "body": "<p>CANARY-HIDDEN-SECTION</p>", "hidden": 1},
    {"id": "s6", "type": "raid", "title": "Risks and issues", "body": "", "hidden": 0},
]
TILES = [
    {"tile": "health", "title": "Project health", "value": "On track", "detail": "", "reason": "Dates hold.", "hidden": False,
     "pill": {"colour": "green"}},
    {"tile": "overdue", "title": "Overdue actions", "value": 3, "detail": "", "hidden": False},
    {"tile": "risks", "title": "High risks and issues", "value": 1, "detail": "of 4 open", "hidden": False},
    {"tile": "decisions", "title": "CANARY-HIDDEN-TILE", "value": 9, "detail": "", "hidden": True},
]
STAMP = "Exported on 01-06-2030 at 20:30 (UTC+04:00)"


def build(**kw):
    args = dict(project=PROJECT, sections=SECTIONS, tiles=TILES, exported=STAMP, by="Sam Example", owner="Kim Example",
                email="kim.example@example.com", internal="INTERNAL: never share this with a client.")
    args.update(kw)
    return cover.lines(**args)


def pin_timezone(case, zone="Asia/Dubai"):
    """Pin the local zone for the test (the exported time is local), and put it back afterwards."""
    old = os.environ.get("TZ")

    def restore():
        if old is None:
            os.environ.pop("TZ", None)
        else:
            os.environ["TZ"] = old
        time.tzset()
    os.environ["TZ"] = zone
    time.tzset()
    case.addCleanup(restore)


def texts(lines):
    return [x["text"] for x in lines]


class TheLines(unittest.TestCase):
    def test_order_title_then_the_sections_then_the_about_block(self):
        got = build()
        self.assertEqual(got[0], {"kind": "title", "text": "Harbour Lights upgrade", "tag": "INTERNAL"})
        headings = [x["text"] for x in got if x["kind"] == "heading"]
        self.assertEqual(headings, ["Current state", "At a glance", "Summary", "This week"])      # the sections' own order
        i = [x["kind"] for x in got].index("about")
        self.assertEqual([x["kind"] for x in got][i:], ["about", "exported", "caution", "return", "where", "ids", "internal"])
        self.assertEqual(got[i]["text"], "About this export")
        self.assertEqual(got[i + 1]["text"], "Exported on 01-06-2030 at 20:30 (UTC+04:00) by Sam Example")
        self.assertEqual(got[-1]["text"], "INTERNAL: never share this with a client.")

    def test_no_internal_note_means_no_tag_and_no_internal_line(self):
        got = build(internal="")
        self.assertNotIn("tag", got[0])
        self.assertNotIn("internal", [x["kind"] for x in got])

    def test_empty_sections_are_left_off_heading_and_all(self):
        secs = [{"type": "current", "title": "Current state"}, {"type": "glance", "title": "At a glance"},
                {"type": "text", "title": "Summary", "body": "<p> </p>"}, {"type": "text", "title": "Next week", "body": "<p>Hand over.</p>"}]
        got = build(project={**PROJECT, "current_state": ""}, sections=secs, tiles=[dict(t, hidden=True) for t in TILES])
        self.assertEqual([x["text"] for x in got if x["kind"] == "heading"], ["Next week"])
        got = build(project={**PROJECT, "current_state": ""}, sections=secs[:2], tiles=[])
        self.assertEqual([x["kind"] for x in got if x["kind"] == "heading"], [])

    def test_the_health_line_carries_the_pills_colour_and_no_other_line_has_one(self):
        got = build()
        self.assertEqual([x for x in got if "colour" in x], [{"kind": "health", "text": "Project health: On track", "colour": "green"}])
        no_colour = [dict(t, pill={"colour": None}) if t["tile"] == "health" else t for t in TILES]
        self.assertEqual([x for x in build(tiles=no_colour) if "colour" in x], [], "health not set: no colour, words only")

    def test_current_state_is_plain_words_with_breaks(self):
        got = texts(build())
        self.assertIn("Phase two is under way.\nCabling done\nTesting next", got)

    def test_the_tiles_and_the_health_reason(self):
        got = texts(build())
        i = got.index("Project health: On track")
        self.assertEqual(got[i + 1], "Dates hold.")
        self.assertIn("Overdue actions: 3", got)
        self.assertIn("High risks and issues: 1 (of 4 open)", got)

    def test_every_text_section_is_there(self):
        got = texts(build())
        self.assertIn("On course for the spring opening.", got)
        self.assertIn("Lights tested on the east pier.", got)

    def test_nothing_hidden_or_internal_reaches_the_cover(self):
        everything = "\n".join(texts(build()))
        for canary in ("CANARY-HIDDEN-SECTION", "CANARY-HIDDEN-TILE", "CANARY-INTERNAL-NOTE", "CANARY-PROJECT-FIELD", "Hidden section"):
            self.assertNotIn(canary, everything)

    def test_a_section_kind_the_cover_does_not_carry_is_left_out(self):
        self.assertNotIn("Risks and issues", [x["text"] for x in build() if x["kind"] == "heading"])

    def test_no_exporter_name_means_just_the_date(self):
        for by in (None, "  "):
            self.assertEqual([x["text"] for x in build(by=by) if x["kind"] == "exported"], [STAMP])


class TheReturnNote(unittest.TestCase):
    def test_with_an_owner_and_an_address(self):
        self.assertIn("send the file back to Kim Example (kim.example@example.com).", cover.return_note("Kim Example", "kim.example@example.com"))

    def test_with_no_address_the_brackets_go(self):
        note = cover.return_note("Kim Example", "")
        self.assertIn("send the file back to Kim Example.", note)
        self.assertNotIn("(", note)

    def test_with_no_name_it_goes_back_to_whoever_sent_it(self):
        self.assertIn("send the file back to the person who sent it to you.", cover.return_note(None, "kim.example@example.com"))
        self.assertNotIn("kim.example", cover.return_note("", "kim.example@example.com"))


class HostileText(unittest.TestCase):
    def test_tags_and_scripts_never_arrive(self):
        project = {**PROJECT, "current_state": "<p>Fine</p><script>alert('x')</script><img src=x onerror=alert(1)><p>=HYPERLINK(\"http://x\")</p>"}
        everything = "\n".join(texts(build(project=project)))
        self.assertNotIn("<", everything)
        self.assertNotIn("alert", everything)
        self.assertIn("Fine", everything)

    def test_a_very_long_body_is_cut_under_excels_limit_with_a_mark(self):
        project = {**PROJECT, "current_state": "<p>" + "word " * 20000 + "</p>"}
        for line in build(project=project):
            self.assertLess(len(line["text"]), 32767)
        longest = max(texts(build(project=project)), key=len)
        self.assertTrue(longest.endswith(cover.MORE))

    def test_a_body_the_cleaner_refuses_is_cut_not_fatal(self):
        project = {**PROJECT, "current_state": "<p>" + "a" * 400000 + "</p>"}
        got = build(project=project)                       # does not raise
        for line in got:
            self.assertLess(len(line["text"]), 32767)

    def test_empty_everything_still_gives_a_cover(self):
        got = cover.lines({}, [], [], STAMP)
        self.assertEqual([x["kind"] for x in got], ["title", "about", "exported", "caution", "return", "where", "ids"])


class PlainWords(unittest.TestCase):
    def test_no_wording_line_starts_with_the_product_name(self):
        for text in (cover.CAUTION, cover.RETURN_WITH, cover.RETURN_ANYONE, cover.WHERE_TO_EDIT, cover.IDS, cover.ABOUT, cover.MORE.strip()):
            self.assertFalse(text.lower().startswith("clark"), text)

    def test_the_wording_says_the_cover_is_not_read_back(self):
        self.assertIn("Changes made on this Cover are ignored", cover.WHERE_TO_EDIT)
        self.assertEqual(cover.IDS, "Please don't change or move IDs. Each ID stays with its own row: if you change one, or sort "
                                    "only the ID column, the details go to another item when the file is imported. Leave the ID "
                                    "empty for a new item.")       # #128 done-when 5, in words true of the import (#129)
        self.assertNotIn("refused", cover.IDS, "#128: a swapped or sorted ID is not refused, so the words must not say it is")


class Escaping(unittest.TestCase):
    """A name, a section title or a tile title that starts like a formula is text on the Cover, never a formula."""
    def setUp(self):
        pin_timezone(self)
        self.dir = Path(tempfile.mkdtemp(prefix="clark-cover-"))
        self.addCleanup(shutil.rmtree, self.dir, True)

    def test_leading_formula_characters_get_an_apostrophe_and_are_never_formulas(self):
        for lead in ("=", "+", "@", "\t"):
            with self.subTest(lead=repr(lead)):
                project = {"name": lead + "Name", "current_state": "<p>Now.</p>"}
                sections = [{"type": "current", "title": lead + "Section"}, {"type": "glance", "title": lead + "Glance"},
                            {"type": "text", "title": lead + "Words", "body": "<p>Body.</p>"}]
                tiles = [{"tile": "overdue", "title": lead + "Tile", "value": 2, "detail": "", "hidden": False}]
                wb = Workbook()
                ws = excel.write_cover(wb, cover.lines(project, sections, tiles, STAMP))
                values = [c.value for row in ws.iter_rows() for c in row if c.value]
                self.assertIn("Name" if lead == "\t" else "'" + lead + "Name", values)      # the name is trimmed, so a tab goes
                for want in ("Section", "Glance", "Words"):
                    self.assertIn("'" + lead + want, values)
                self.assertIn("'" + lead + "Tile: 2", values)
                for row in ws.iter_rows():
                    for c in row:
                        self.assertNotEqual(c.data_type, "f", c.coordinate)


class InTheWorkbook(unittest.TestCase):
    """The Cover sheet through the real export of a demo store: where it sits, what it holds, what it must not, and that the
    import never reads it."""
    ON = "2030-06-01"

    def setUp(self):
        pin_timezone(self)
        self.demo = DemoStore()
        self.s = self.demo.store
        self.dir = Path(tempfile.mkdtemp(prefix="clark-cover-"))
        self.addCleanup(shutil.rmtree, self.dir, True)
        self.addCleanup(self.demo.close)
        self.path, self.state = self.dir / "ABC work items.xlsx", self.dir / "state.json"
        self.demo.action(description="Visible overdue action", due_date="2030-01-01")
        self.demo.update("milestones", self.demo.milestone, {"internal_note": "CANARY-INTERNAL-NOTE"})
        self.demo.action(description="CANARY-SENSITIVE-ROW", due_date="2030-01-01", labels="Sensitive")
        self.demo.action(description="CANARY-NOT-FOR-CLIENT", due_date="2030-01-01", client_visible=0)
        con = sqlite3.connect(self.s.path)
        con.execute("UPDATE projects SET current_state = ? WHERE code = 'ABC'",
                    ("<p>=HYPERLINK(\"http://x\")</p><script>alert(1)</script><p>Phase two is under way.</p>",))
        con.execute("DELETE FROM report_sections WHERE project = 'ABC'")      # the default sections a new project gets (#65): only these four
        con.execute("INSERT INTO report_sections (project, id, type, title, body, beside, hidden, sort_order) VALUES "
                    "('ABC','a','current','Current state','',0,0,1), ('ABC','b','glance','At a glance','',0,0,2), "
                    "('ABC','c','text','Summary','<p>On course.</p>',0,0,3), "
                    "('ABC','d','text','Secret plans','<p>CANARY-HIDDEN-SECTION</p>',0,1,4)")
        con.commit()
        con.close()
        self.make()

    def make(self):
        return excel.export_project(self.s, "ABC", self.path, self.state, on=self.ON, exported_by="Sam Example", owner="Kim Example",
                                    owner_email="kim.example@example.com")

    def cover_cells(self):
        ws = load_workbook(self.path)["Cover"]
        return [c.value for row in ws.iter_rows() for c in row if c.value is not None]

    def test_cover_is_first_and_open_with_no_merged_cells_and_the_top_row_frozen(self):
        wb = load_workbook(self.path)
        self.assertEqual(wb.sheetnames, ["Cover", "Milestones", "Actions", "RAID", "Decisions", "Use cases", "Lists"])
        self.assertEqual(wb.active.title, "Cover")
        self.assertEqual(list(wb["Cover"].merged_cells.ranges), [])
        self.assertEqual(wb["Cover"].freeze_panes, "A2")      # below the title band

    def cover_sheet(self):
        return load_workbook(self.path)["Cover"]

    def find(self, value):
        (cell,) = [c for row in self.cover_sheet().iter_rows() for c in row if c.value == value]
        return cell

    def test_the_title_band_is_navy_with_white_bold_text_and_the_internal_tag(self):
        ws = self.cover_sheet()
        a, b = ws["A1"], ws["B1"]
        self.assertEqual((a.value, b.value), ("INTERNAL", "Example Programme"))
        self.assertEqual(a.fill.fgColor.rgb[-6:], colours.TOKENS["status-red"], "the INTERNAL tag is red")
        self.assertEqual(b.fill.fgColor.rgb[-6:], colours.TOKENS["brand-navy"])
        for c in (a, b):
            self.assertEqual(c.font.color.rgb[-6:], "FFFFFF")
            self.assertTrue(c.font.b)

    def test_section_headings_have_the_brand_blue_marker_and_tiles_are_quiet_grey(self):
        h = self.find("Summary")
        self.assertEqual(h.font.color.rgb[-6:], colours.TOKENS["brand-blue"])
        self.assertTrue(h.font.b)
        self.assertEqual(h.border.left.color.rgb[-6:], colours.TOKENS["brand-blue"])
        tile = self.find("Overdue actions: 1 (of 1 open action)")
        self.assertEqual(tile.fill.fgColor.rgb[-6:], colours.TOKENS["grey-50"])

    def test_the_cover_reads_in_the_dashboards_order_and_ends_with_the_about_block(self):
        col = [c.value for c in self.cover_sheet()["A"] if c.value]
        self.assertEqual(col[:4], ["INTERNAL", "Current state", "At a glance", "Summary"])
        tail = [c.value for c in self.cover_sheet()["A"]][-6:]
        self.assertEqual(tail, [None, "Caution", "Sending changes back", "Where to make changes", "IDs", "Internal"])

    def test_the_about_block_is_quiet_grey(self):
        head = self.find("About this export")
        self.assertEqual(head.fill.fgColor.rgb[-6:], colours.TOKENS["grey-50"])
        self.assertEqual(head.font.color.rgb[-6:], colours.TOKENS["grey-700"])

    def test_the_word_provenance_is_nowhere(self):
        ws = self.cover_sheet()
        self.assertNotIn("provenance", "\n".join(str(c.value) for row in ws.iter_rows() for c in row).lower())
        self.assertNotIn("Provenance", repr(excel.COVER_LABELS))

    def test_everything_on_the_cover_is_arial_text_and_no_line_starts_with_the_product_name(self):
        ws = self.cover_sheet()
        for row in ws.iter_rows():
            for c in row:
                self.assertEqual(c.font.name, "Arial", c.coordinate)
                self.assertIn(c.data_type, ("s", "inlineStr"), c.coordinate)
                self.assertFalse(str(c.value or "").lower().startswith("clark"), c.coordinate)

    def test_what_it_says(self):
        text = self.cover_cells()
        self.assertIn("Example Programme", text)
        self.assertTrue(any(c.startswith("INTERNAL:") for c in text))
        stamp = [c for c in text if c.startswith("Exported ")]
        self.assertEqual(len(stamp), 1)
        self.assertRegex(stamp[0], r"^Exported on \d{2}-\d{2}-\d{4} at ([01]\d|2[0-3]):[0-5]\d \(UTC[+-]\d{2}:\d{2}\) by Sam Example$")
        self.assertIn(cover.CAUTION, text)
        self.assertTrue(any("Kim Example (kim.example@example.com)" in c for c in text))
        self.assertIn("On course.", text)
        self.assertIn("Project health: On track", text)

    def test_counts_are_the_pages_counts_so_hidden_and_sensitive_rows_are_not_in_them(self):
        self.assertIn("Overdue actions: 1 (of 1 open action)", self.cover_cells())

    def test_no_canary_reaches_the_cover(self):
        everything = "\n".join(self.cover_cells())
        for canary in ("CANARY-INTERNAL-NOTE", "CANARY-SENSITIVE-ROW", "CANARY-NOT-FOR-CLIENT", "CANARY-HIDDEN-SECTION", "Secret plans"):
            self.assertNotIn(canary, everything)
        data = "\n".join(str(c.value) for sheet in ("Milestones", "Actions") for row in load_workbook(self.path)[sheet].iter_rows() for c in row)
        for canary in ("CANARY-INTERNAL-NOTE", "CANARY-SENSITIVE-ROW", "CANARY-NOT-FOR-CLIENT"):
            self.assertIn(canary, data, "the data sheets still hold them: only the Cover is clean")

    def test_a_hostile_current_state_makes_no_tag_and_no_formula(self):
        wb = load_workbook(self.path)
        for row in wb["Cover"].iter_rows():
            for c in row:
                self.assertNotEqual(c.data_type, "f", c.coordinate)
        everything = "\n".join(self.cover_cells())
        self.assertNotIn("<script", everything)
        self.assertNotIn("alert", everything)
        self.assertIn("'=HYPERLINK(\"http://x\")\nPhase two is under way.", self.cover_cells(),
                      "a cell that starts with = gets the apostrophe; a later line in it can't start a formula")

    def test_the_health_pill_is_a_coloured_cell_and_its_words_name_the_health(self):
        ws = load_workbook(self.path)["Cover"]
        (cell,) = [c for row in ws.iter_rows() for c in row if c.value == "Project health: On track"]
        fill, ink = colours.hexes("green")
        self.assertEqual(cell.fill.fgColor.rgb[-6:], fill)
        self.assertEqual(cell.font.color.rgb[-6:], ink)
        others = [c for row in ws.iter_rows() for c in row if c.value and str(c.value).startswith("Overdue actions")]
        self.assertTrue(others and all(c.fill.fgColor.rgb[-6:] == colours.TOKENS["grey-50"] for c in others),
                        "the other tiles are quiet grey; only the pill has a status colour")

    def test_the_import_never_reads_the_cover(self):
        plan = xi.read_workbook(self.s, self.path, "ABC", state_path=self.state)
        self.assertEqual((plan.changes, plan.problems), ([], []))
        wb = load_workbook(self.path)
        for row in wb["Cover"].iter_rows():
            for c in row:
                if c.value is not None:
                    c.value = "edited " + str(c.value)
        wb.save(self.path)
        plan = xi.read_workbook(self.s, self.path, "ABC", state_path=self.state)
        self.assertEqual((plan.changes, plan.problems), ([], []), "a Cover with every cell edited changes nothing")
        self.assertEqual([n for n in plan.notes if "ignored" in n], [])

    def test_the_same_data_and_clock_give_the_same_bytes(self):
        first = self.path.read_bytes()
        self.state.unlink()
        self.path.unlink()
        self.make()
        self.assertEqual(self.path.read_bytes(), first)


if __name__ == "__main__":
    unittest.main()
