"""#444 (rows 44): Settings › Data in a real browser against the real routes: Import workbook (the two-step box, one step, Undo this import)
and Reset the demo data (offered only for demo data, asks first, says so when it fails). Run: python3 dev/run_tests.py test_settings_data_page

The server is the real one over the #39 plant in a temporary folder (tests/test_publish_route.py Rig) with the real workbook import and the
real demo seam; the demo seam's rebuild is a recorder, so nothing is replaced. Widths 1400 and 900, light and dark. Chrome is required:
these tests fail, never skip, without it. The non-browser checks (the published page names none of these routes) need no browser.
"""
import base64
import json
import tempfile
import unittest
from pathlib import Path

from _paths import SRC
from test_publish_dialog import Base
from test_server import chrome, wait

from clark import demo_route, excel, excel_route

WIDTHS = (1400, 900)
USER = "Test Person"


def workbook(rig, tmp, new_text):
    """A workbook exported as the page's Export does it, with the first action's description changed; (base64, the action's id, its old text)."""
    from openpyxl import load_workbook
    path = Path(tmp) / "PUB work items.xlsx"
    made = excel.export_project(rig.store, "PUB", path, Path(tmp) / "state.json", on="2030-06-01")
    record = excel.load_state(Path(tmp) / "state.json", made)
    excel.save_state(rig.data_dir / excel.STATE_FILE, excel.id_key("PUB", record["export_id"]), record)
    action = rig.store.read("actions", "PUB")[0]
    wb = load_workbook(str(path))
    ws = wb["Actions"]
    head = {c.value: c.column for c in ws[1]}
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, head["ID"]).value == action["id"])
    ws.cell(row, head["Description"]).value = new_text
    wb.save(str(path))
    return base64.b64encode(path.read_bytes()).decode(), action["id"], action["description"]


class Source(unittest.TestCase):
    def test_the_published_page_names_none_of_these_routes_or_boxes(self):
        view = (SRC / "clark" / "page" / "page-view.html").read_text(encoding="utf-8")
        for needle in ("/api/excel-import", "/api/demo", "Import workbook", "Reset the demo data", "import-box", "reset-box"):
            self.assertNotIn(needle, view, needle)

    def test_only_the_data_files_name_the_three_routes(self):
        found = {}
        for f in (SRC / "clark" / "page" / "src").rglob("*"):
            if f.is_file() and f.suffix in (".js", ".jsx"):
                text = f.read_text(encoding="utf-8")
                for route in ('"/api/excel-import"', '"/api/demo"', '"/api/project"'):
                    if route in text:
                        found.setdefault(route, []).append(f.name)
        self.assertEqual(found, {'"/api/excel-import"': ["dataApi.js"], '"/api/demo"': ["dataApi.js"], '"/api/project"': ["dataApi.js"]})


class Card(Base):
    def setUp(self):
        super().setUp()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.rig.httpd.excel_import = excel_route.make_excel_import(self.rig.db, [], USER)
        self.rebuilt = []
        self.fails = False

        def rebuild(path):
            if self.fails:
                raise RuntimeError("no")
            self.rebuilt.append(str(path))
        self.rig.httpd.demo = demo_route.make_demo(self.rig.db, self.rig.data_dir, rebuild)

    def settle(self, p, expr, what, seconds=20):
        def got():
            try:
                return bool(p.js(expr))
            except Exception:                    # the page is reloading
                return False
        self.assertTrue(wait(got, seconds), what)

    def open_settings(self, p):
        p.js("document.querySelector('#profile-btn').click()")
        p.until("!!document.querySelector('#profile-menu')", "the menu opened")
        p.js("document.querySelector('#profile-menu [data-kind=settings]').click()")
        p.until("!!document.querySelector('#data-card')", "the Data card drew")

    def choose(self, p, b64):
        p.js("""(() => { const bin = atob(%s), a = new Uint8Array(bin.length); for (let i = 0; i < bin.length; i++) a[i] = bin.charCodeAt(i);
          const dt = new DataTransfer(); dt.items.add(new File([a], 'PUB work items.xlsx')); const inp = document.querySelector('#import-file');
          inp.files = dt.files; inp.dispatchEvent(new Event('change', {bubbles: true})); })()""" % json.dumps(b64))

    def test_the_card_offers_reset_only_for_demo_data_with_the_prototypes_three_buttons(self):
        for w in WIDTHS:
            with self.subTest(width=w), chrome() as cdp:
                p = self.page(cdp)
                p.width(w)
                self.open_settings(p)
                self.settle(p, "document.querySelector('#reset-demo')", "Reset the demo data on the card")
                self.assertEqual(p.js("[...document.querySelectorAll('#data-card .row button')].map(b => b.textContent)"),
                                 ["Export workbook", "Import workbook", "Reset the demo data"])
                self.assertLessEqual(p.js("document.documentElement.scrollWidth"), w)

    def test_import_checks_then_imports_in_one_step_and_undo_puts_it_back(self):
        b64, action, old = workbook(self.rig, self.tmp.name, "Agree the cut-over plan, version two")
        for w in WIDTHS:
            with self.subTest(width=w), chrome() as cdp:
                p = self.page(cdp)
                p.width(w)
                self.open_settings(p)
                p.js("document.querySelector('#import-workbook').click()")
                p.until("!!document.querySelector('#import-box')", "the box opened")
                self.assertEqual(p.js("document.querySelector('#import-box-t').textContent"), "Import from Excel")
                self.assertEqual(p.js("document.activeElement.id"), "import-choose")
                self.choose(p, b64)
                p.until("!!document.querySelector('#import-summary')", "the check")
                self.assertEqual(p.js("document.querySelector('#import-box-t').textContent"), "Check the import")
                self.assertIn("1 changed", p.js("document.querySelector('#import-summary').textContent"))
                self.assertIn(action, p.js("document.querySelector('.imp-list li').innerText"))
                self.assertEqual(self.rig.store.read("actions", ref=action)["description"], old, "a check writes nothing")
                self.assertLessEqual(p.js("document.documentElement.scrollWidth"), w)
                steps = self.rig.store.con.execute("SELECT count(*) FROM applies").fetchone()[0]
                p.js("document.querySelector('#import-go').click()")
                p.until("!!document.querySelector('#import-result')", "the result")
                self.assertIn("1 changed", p.js("document.querySelector('#import-result').textContent"))
                self.assertEqual(self.rig.store.read("actions", ref=action)["description"], "Agree the cut-over plan, version two")
                self.assertEqual(self.rig.store.con.execute("SELECT count(*) FROM applies").fetchone()[0], steps + 1, "ONE step")
                p.js("document.querySelector('#import-undo').click()")
                p.until("document.querySelector('#import-result') && document.querySelector('#import-result').textContent.startsWith('Put back')", "Undo")
                self.assertEqual(self.rig.store.read("actions", ref=action)["description"], old)
                p.js("document.querySelector('#import-close').click()")
                p.until("!document.querySelector('#import-box')", "the box closed")

    def test_a_file_that_is_not_a_workbook_is_said_in_words_in_the_first_box(self):
        with chrome() as cdp:
            p = self.page(cdp)
            self.open_settings(p)
            p.js("document.querySelector('#import-workbook').click()")
            p.until("!!document.querySelector('#import-box')", "the box opened")
            self.choose(p, base64.b64encode(b"not a workbook").decode())
            p.until("!!document.querySelector('#import-problem').textContent", "the refusal")
            self.assertIn("can't be read as an Excel workbook", p.js("document.querySelector('#import-problem').textContent"))
            self.assertTrue(p.js("!!document.querySelector('#import-choose') && !document.querySelector('#import-summary')"))

    def test_escape_closes_the_box_and_focus_goes_back_to_the_button(self):
        with chrome() as cdp:
            p = self.page(cdp)
            self.open_settings(p)
            p.js("document.querySelector('#import-workbook').click()")
            p.until("!!document.querySelector('#import-box')", "the box opened")
            p.js("document.querySelector('#import-box').dispatchEvent(new KeyboardEvent('keydown', {key: 'Escape', bubbles: true}))")
            p.until("!document.querySelector('#import-box')", "Esc closed it")
            self.settle(p, "document.activeElement && document.activeElement.id === 'import-workbook'", "focus back on Import workbook")

    def test_reset_asks_first_cancel_changes_nothing_and_reset_rebuilds_and_says_so(self):
        with chrome() as cdp:
            p = self.page(cdp)
            self.open_settings(p)
            self.settle(p, "document.querySelector('#reset-demo')", "the Reset button")
            p.js("document.querySelector('#reset-demo').click()")
            p.until("!!document.querySelector('#reset-box')", "the confirm")
            self.assertEqual(p.js("document.querySelector('#reset-box-t').textContent"), "Reset the demo data?")
            self.assertIn("Every change you made here is replaced by the starting demo data.", p.js("document.querySelector('#reset-box').innerText"))
            self.assertEqual(p.js("document.activeElement.id"), "reset-cancel", "Cancel has focus: Enter does not reset")
            p.js("document.querySelector('#reset-cancel').click()")
            p.until("!document.querySelector('#reset-box')", "Cancel closed it")
            self.assertEqual(self.rebuilt, [])
            p.js("document.querySelector('#reset-demo').click()")
            p.until("!!document.querySelector('#reset-go')", "the confirm again")
            p.js("document.querySelector('#reset-go').click()")
            self.settle(p, "document.querySelector('#reset-done') && document.querySelector('#reset-done').textContent === 'The demo data is back as it started.'", "the page reloaded and said so")
            self.assertEqual(self.rebuilt, [str(self.rig.db)])

    def test_a_reset_that_fails_says_so_and_the_box_stays(self):
        self.fails = True
        with chrome() as cdp:
            p = self.page(cdp)
            self.open_settings(p)
            self.settle(p, "document.querySelector('#reset-demo')", "the Reset button")
            p.js("document.querySelector('#reset-demo').click()")
            p.until("!!document.querySelector('#reset-go')", "the confirm")
            p.js("document.querySelector('#reset-go').click()")
            p.until("!!document.querySelector('#reset-problem').textContent", "the refusal")
            self.assertEqual(p.js("document.querySelector('#reset-problem').textContent"), demo_route.FAILED)
            self.assertTrue(p.js("!!document.querySelector('#reset-box') && !document.querySelector('#reset-go').disabled"))

    def test_the_buttons_and_the_boxes_are_readable_in_light_and_dark(self):
        b64, _, _ = workbook(self.rig, self.tmp.name, "Agree the cut-over plan, again")
        probe = """(sel) => { const b = document.querySelector(sel); const cs = getComputedStyle(b);
          const px = (c) => c.match(/[\\d.]+/g).map(Number).slice(0, 3);
          const lum = ([r, g, bl]) => { const f = (v) => { v /= 255; return v <= .03928 ? v / 12.92 : Math.pow((v + .055) / 1.055, 2.4); }; return .2126 * f(r) + .7152 * f(g) + .0722 * f(bl); };
          let bg = b, c; while (bg) { c = getComputedStyle(bg).backgroundColor; if (!/rgba\\(.*, 0\\)|transparent/.test(c)) break; bg = bg.parentElement; }
          const a = lum(px(cs.color)), z = lum(px(c)); return (Math.max(a, z) + .05) / (Math.min(a, z) + .05); }"""
        for scheme in ("light", "dark"):
            with self.subTest(scheme=scheme), chrome() as cdp:
                p = self.page(cdp)
                p.js(f"document.documentElement.setAttribute('data-theme', '{scheme}')")
                self.open_settings(p)
                self.settle(p, "document.querySelector('#reset-demo')", "the Reset button")
                for sel in ("#export-workbook", "#import-workbook", "#reset-demo"):
                    self.assertGreaterEqual(p.js(f"({probe})('{sel}')"), 4.5, sel)
                p.js("document.querySelector('#import-workbook').click()")
                p.until("!!document.querySelector('#import-box')", "the box opened")
                self.choose(p, b64)
                p.until("!!document.querySelector('#import-go')", "the check")
                for sel in ("#import-go", "#import-box .btn.quiet"):
                    self.assertGreaterEqual(p.js(f"({probe})('{sel}')"), 4.5, sel)


if __name__ == "__main__":
    unittest.main()
