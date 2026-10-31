"""#133: the brand logo on the Excel Cover, put in after openpyxl has saved, with no Pillow (clark.excel_picture). The logo is the
user's own (system/branding, not in the repo), so every test makes its own tiny PNG in a temporary brand folder.

Run: python3 dev/run_tests.py test_excel_picture
openpyxl reads a workbook back without Pillow by dropping its pictures, so these tests look in the saved zip.
"""
import shutil
import struct
import sqlite3
import tempfile
import unittest
import zipfile
import zlib
from pathlib import Path
from xml.etree import ElementTree as ET

import _paths  # noqa: F401  (puts src/ and vendor/ on the path)
from _demo import DemoStore
from openpyxl import load_workbook

from clark import excel, excel_picture as pic
from clark import excel_cover as cover
from clark import excel_import as xi

PROJECT = {"code": "DEMO", "name": "Harbour Lights upgrade", "current_state": "<p>Under way.</p>", "internal_notes": "", "health_reason": ""}
STAMP = "Exported 01-06-2030 20:30"
XDR = "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}"


def chunk(kind, body):
    return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body))


def make_png(w=3, h=2, first=b"IHDR"):
    return (pic.SIGNATURE + chunk(first, struct.pack(">IIBBBBB", w, h, 8, 6, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(b"\0" * (w * 4 * h + h)))
            + chunk(b"IEND", b""))


class Reader(unittest.TestCase):
    def test_sizes_come_from_the_header(self):
        self.assertEqual(pic.read_size(make_png(3, 2)), (3, 2))
        self.assertEqual(pic.read_size(make_png(2000, 1)), (2000, 1))

    def test_refusals_are_plain(self):
        good = make_png()
        cases = {"not a png": b"GIF89a" + b"\0" * 40, "empty": b"", "truncated": good[:20], "signature only": pic.SIGNATURE,
                 "huge picture": make_png(2001, 5), "zero wide": make_png(0, 5), "other chunk first": make_png(first=b"IDAT"),
                 "too many bytes": good + b"\0" * pic.MAX_BYTES,
                 "a damaged chunk": good[:40] + bytes([good[40] ^ 1]) + good[41:],
                 "cut after the header": good[:33],
                 "no end chunk": good[:-12],
                 "bytes after the end": good + b"extra"}
        for label, data in cases.items():
            with self.subTest(label), self.assertRaises(pic.PictureRefused) as cm:
                pic.read_size(data, "logo.png")
            self.assertIn("logo.png", str(cm.exception))


class Lookup(unittest.TestCase):
    def setUp(self):
        self.brand = Path(tempfile.mkdtemp(prefix="clark-brand-"))
        self.addCleanup(shutil.rmtree, self.brand, True)

    def test_a_logo_in_the_brand_folder_is_found(self):
        (self.brand / "logo.png").write_bytes(make_png(8, 4))
        self.assertEqual(pic.local_logo(self.brand), self.brand / "logo.png")
        self.assertEqual(pic.load(self.brand), (make_png(8, 4), 8, 4))

    def test_a_logo_in_the_design_system_folder_is_found(self):
        (self.brand / "design-system").mkdir()
        (self.brand / "design-system" / "logo.png").write_bytes(make_png())
        self.assertEqual(pic.local_logo(self.brand), self.brand / "design-system" / "logo.png")

    def test_no_logo_is_none(self):
        self.assertIsNone(pic.local_logo(self.brand))
        self.assertIsNone(pic.load(self.brand))

    def test_the_default_folder_is_the_pages_local_brand(self):
        from clark import page
        old = page.BRANDING
        page.BRANDING = self.brand
        self.addCleanup(setattr, page, "BRANDING", old)
        (self.brand / "logo.png").write_bytes(make_png())
        self.assertEqual(pic.local_logo(), self.brand / "logo.png")

    def test_a_bad_logo_gives_none_and_a_plain_warning(self):
        (self.brand / "logo.png").write_bytes(b"not a picture")
        with self.assertLogs("clark.excel", "WARNING") as logs:
            self.assertIsNone(pic.load(self.brand))
        self.assertIn("logo.png is not a PNG picture", logs.output[0])
        self.assertIn("The brand logo was not used, so the Cover has no logo", logs.output[0])

    def test_the_reason_is_handed_back_so_the_screen_can_say_why_the_logo_is_missing(self):
        """"""
        (self.brand / "logo.png").write_bytes(make_png()[:-3])
        said = []
        with self.assertLogs("clark.excel", "WARNING"):
            self.assertIsNone(pic.load(self.brand, said))
        self.assertEqual(said, ["The brand logo was not used, so the Cover has no logo: logo.png is cut short: it ends in the "
                                "middle of the picture."])
        self.assertIsNone(pic.load(Path(tempfile.mkdtemp(prefix="clark-none-")), said), "no logo at all is not a warning")
        self.assertEqual(len(said), 1)

    def test_the_read_is_bounded(self):
        """"""
        (self.brand / "logo.png").write_bytes(make_png() + b"\0" * (2 * pic.MAX_BYTES))
        real_open = open
        seen = []

        class Watch:
            def __init__(self, f):
                self.f = f

            def __enter__(self):
                return self

            def __exit__(self, *a):
                self.f.close()

            def read(self, n=-1):
                seen.append(n)
                return self.f.read(n)

        from unittest import mock
        with mock.patch("builtins.open", lambda p, m="r", *a, **k: Watch(real_open(p, m, *a, **k))), \
                self.assertLogs("clark.excel", "WARNING"):
            self.assertIsNone(pic.load(self.brand))
        self.assertEqual(seen, [pic.MAX_BYTES + 1])


class InTheWorkbook(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp(prefix="clark-mark-"))
        self.addCleanup(shutil.rmtree, self.dir, True)
        self.brand = self.dir / "branding"
        self.brand.mkdir()
        self.png = make_png(60, 20)
        self.path, self.state = self.dir / "w.xlsx", self.dir / "state.json"

    def export(self, name="w.xlsx"):
        self.path = self.dir / name
        lines = cover.lines(project=PROJECT, sections=[], tiles=[], exported=STAMP, by="Sam Example", owner="Kim Example",
                            email="kim.example@example.com", internal="INTERNAL: never share this with a client.")
        excel.export(self.path, "DEMO", {}, self.state, on="2030-06-01", cover=lines, branding=self.brand)
        return zipfile.ZipFile(self.path)

    def parts(self, z):
        return [n for n in z.namelist() if "media" in n or "drawings/drawing" in n]

    def test_a_generated_logo_is_embedded_byte_for_byte_on_the_cover(self):
        (self.brand / "logo.png").write_bytes(self.png)
        z = self.export()
        self.assertIsNone(z.testzip())
        self.assertEqual(z.read("xl/media/image1.png"), self.png)
        self.assertEqual(self.parts(z), ["xl/drawings/drawing1.xml", "xl/media/image1.png"])
        # The Cover is sheet1, and it points at the drawing, which points at the image.
        self.assertIn(b'Target="../drawings/drawing1.xml"', z.read("xl/worksheets/_rels/sheet1.xml.rels"))
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
        self.assertEqual(sheet.count("<drawing "), 1)
        self.assertNotIn("<drawing ", z.read("xl/worksheets/sheet2.xml").decode())
        rid = ET.fromstring(z.read("xl/worksheets/sheet1.xml")).find(
            "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}drawing").get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        self.assertIn(f'Id="{rid}"', z.read("xl/worksheets/_rels/sheet1.xml.rels").decode())
        self.assertIn(b'Target="../media/image1.png"', z.read("xl/drawings/_rels/drawing1.xml.rels"))
        types = z.read("[Content_Types].xml").decode()
        self.assertIn('<Default Extension="png" ContentType="image/png"/>', types)
        self.assertIn('PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"', types)
        ET.fromstring(types)

    def test_the_drawing_anchors_at_the_right_of_the_band_scaled_to_the_row(self):
        (self.brand / "logo.png").write_bytes(self.png)
        root = ET.fromstring(self.export().read("xl/drawings/drawing1.xml"))
        anchor = root.find(XDR + "oneCellAnchor")
        self.assertEqual(anchor.find(XDR + "from/" + XDR + "col").text, "1")
        self.assertEqual(anchor.find(XDR + "from/" + XDR + "row").text, "0")
        ext = anchor.find(XDR + "ext")
        self.assertEqual((int(ext.get("cx")), int(ext.get("cy"))), (120 * 9525, 40 * 9525))
        left = int(anchor.find(XDR + "from/" + XDR + "colOff").text)
        self.assertEqual(left + 120 * 9525 + pic.MARK_MARGIN * 9525, pic.B_WIDTH_PIXELS * 9525)
        self.assertEqual(load_workbook(self.path)["Cover"].row_dimensions[1].height, pic.COVER_ROW_POINTS)

    def test_the_workbook_still_loads(self):
        (self.brand / "logo.png").write_bytes(self.png)
        self.export()
        wb = load_workbook(self.path)
        self.assertEqual(wb.sheetnames[0], "Cover")
        self.assertEqual(wb["Cover"]["B1"].fill.fgColor.rgb[-6:], excel.colours.TOKENS["brand-navy"])

    def test_the_same_data_and_clock_give_the_same_bytes_and_dates(self):
        (self.brand / "logo.png").write_bytes(self.png)
        z = self.export()
        first = self.path.read_bytes()
        self.assertEqual({i.date_time for i in z.infolist()}, {excel.ZIP_TIME})
        self.assertEqual({i.compress_type for i in z.infolist()}, {zipfile.ZIP_DEFLATED})
        self.assertEqual(z.namelist(), sorted(z.namelist()))
        z.close()
        self.state.unlink()
        self.path.unlink()
        self.export()
        self.assertEqual(self.path.read_bytes(), first)

    def test_no_logo_means_no_drawing_and_no_picture_row_height(self):
        z = self.export()
        self.assertEqual(self.parts(z), [])
        self.assertNotIn("<drawing ", z.read("xl/worksheets/sheet1.xml").decode())
        self.assertNotIn("image/png", z.read("[Content_Types].xml").decode())
        self.assertIsNone(load_workbook(self.path)["Cover"].row_dimensions[1].height)

    def test_a_bad_logo_is_refused_in_words_and_the_export_goes_ahead_without_it(self):
        (self.brand / "logo.png").write_bytes(make_png(3000, 10))
        with self.assertLogs("clark.excel", "WARNING") as logs:
            z = self.export()
        self.assertIn("3000 by 10 pixels", logs.output[0])
        self.assertEqual(self.parts(z), [])
        self.assertTrue(self.path.exists())
        self.assertEqual(load_workbook(self.path).sheetnames[0], "Cover")

    def test_a_sheet_with_its_own_relationships_keeps_them(self):
        (self.brand / "logo.png").write_bytes(self.png)
        self.export()
        with zipfile.ZipFile(self.path) as z:
            files = {n: z.read(n) for n in z.namelist() if n != "xl/worksheets/_rels/sheet1.xml.rels"}
            files["xl/worksheets/_rels/sheet1.xml.rels"] = (
                b'<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                b'<Relationship Id="rId1" Type="x" Target="y"/></Relationships>')
            for n in ("xl/drawings/drawing1.xml", "xl/drawings/_rels/drawing1.xml.rels", "xl/media/image1.png"):
                del files[n]
            files["xl/worksheets/sheet1.xml"] = files["xl/worksheets/sheet1.xml"].replace(b"<drawing ", b"<zz ").split(b"<zz ")[0] + b"</worksheet>"
            files["[Content_Types].xml"] = files["[Content_Types].xml"].replace(
                b'<Default Extension="png" ContentType="image/png"/>', b"")
        replace, add = pic.embed(files, self.png)
        self.assertIn(b'Id="rId1"', replace["xl/worksheets/_rels/sheet1.xml.rels"])
        self.assertIn(b'Id="rId2"', replace["xl/worksheets/_rels/sheet1.xml.rels"])
        self.assertIn(b'Extension="png"', replace["[Content_Types].xml"])

    def test_a_workbook_with_no_cover_refuses_in_words(self):
        with self.assertRaises(pic.PictureRefused):
            pic.embed({"xl/workbook.xml": b'<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheets/></workbook>',
                       "xl/_rels/workbook.xml.rels": b"<Relationships/>"}, self.png)

    def test_the_write_is_atomic_and_leaves_no_temporary_file(self):
        (self.brand / "logo.png").write_bytes(self.png)
        self.export()
        self.assertEqual(sorted(p.name for p in self.dir.iterdir()), ["branding", "state.json", "w.xlsx"])


class WithAStore(unittest.TestCase):
    """The real export of a demo store: the import gives no changes with a logo on the Cover."""

    def test_the_import_ignores_the_cover_with_its_logo(self):
        demo = DemoStore()
        self.addCleanup(demo.close)
        d = Path(tempfile.mkdtemp(prefix="clark-mark-store-"))
        self.addCleanup(shutil.rmtree, d, True)
        (d / "branding").mkdir()
        (d / "branding" / "logo.png").write_bytes(make_png(30, 30))
        path, state = d / "ABC work items.xlsx", d / "state.json"
        excel.export_project(demo.store, "ABC", path, state, on="2030-06-01", exported_by="Sam Example", branding=d / "branding")
        with zipfile.ZipFile(path) as z:
            self.assertIn("xl/media/image1.png", z.namelist())
        plan = xi.read_workbook(demo.store, path, "ABC", state_path=state)
        self.assertEqual((plan.changes, plan.problems), ([], []))


if __name__ == "__main__":
    unittest.main()
